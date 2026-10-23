# -*- coding: utf-8 -*-
"""
excel_dashboard.py - the purchase-orders dashboard as ONE Excel page that looks like a management report:
title banner, four KPI cards, five native Excel charts in the dashboard colours, ready to print on A4 landscape.

The charts are real Excel charts (click one and «Edit data» works), fed from the «البيانات» sheet, which holds
exactly the numbers shown on the screen (they come from analytics_service, the same code as the web page).
"""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY, TEAL, SAGE, GOLD, RUST = "1E293B", "29595E", "A4BFC1", "C08D34", "B9695C"
PALETTE = ["29595E", "A4BFC1", "C08D34", "6F9EA3", "E0B866", "8A8F98"]
STATUS_COLORS = {"مكتمل": TEAL, "بانتظار السداد": GOLD, "بانتظار الفاتورة": SAGE, "بانتظار الاستلام": RUST}
PAGE_BG, WHITE, LINE, MUTED, TEXT = "F4F6F8", "FFFFFF", "E5EAF0", "64748B", "0F172A"
GOOD, GOOD_BG, BAD, BAD_BG = "15803D", "DCFCE7", "B91C1C", "FEE2E2"
FONT = "Segoe UI"
COLS = 16
COL_W = 11.6


def _fill(hex_color):
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def _font(size=11, bold=False, color=TEXT):
    return Font(name=FONT, size=size, bold=bold, color=color)


def _box(ws, r1, c1, r2, c2, fill=WHITE, border=LINE):
    side = Side(style="thin", color=border)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            cell.fill = _fill(fill)
            cell.border = Border(left=side if c == c1 else None, right=side if c == c2 else None,
                                 top=side if r == r1 else None, bottom=side if r == r2 else None)


def _merge(ws, r1, c1, r2, c2, value=None, font=None, align="center", fill=None, fmt=None, valign="center"):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    cell = ws.cell(r1, c1)
    if value is not None:
        cell.value = value
    if font:
        cell.font = font
    cell.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=True, readingOrder=2)
    if fill:
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(r, c).fill = _fill(fill)
    if fmt:
        cell.number_format = fmt
    return cell


def _style_chart(ch, title=None):
    ch.title = None
    ch.legend = None
    ch.style = 2
    ch.roundedCorners = False
    ch.graphical_properties = GraphicalProperties(solidFill=WHITE, ln=LineProperties(noFill=True))
    ch.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    return ch


def _axes(ch, y_fmt="#,##0", show_y=True, grid=True):
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.number_format = y_fmt
    ch.y_axis.majorGridlines = ch.y_axis.majorGridlines if grid else None
    if grid and ch.y_axis.majorGridlines is not None:
        ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE, w=6350))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=LINE))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    if not show_y:
        ch.y_axis.delete = True


def _labels(series, fmt="#,##0", pos=None, size=900):
    series.dLbls = DataLabelList()
    series.dLbls.showVal = True
    for flag in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
        setattr(series.dLbls, flag, False)
    series.dLbls.numFmt = fmt
    if pos:
        series.dLbls.position = pos


def _panel(ws, r1, c1, r2, c2, title, subtitle):
    _box(ws, r1, c1, r2, c2)
    _merge(ws, r1 + 1, c1, r1 + 1, c2, title, _font(13, True), align="right").alignment = Alignment(horizontal="right", vertical="center", indent=1, readingOrder=2)
    _merge(ws, r1 + 2, c1, r1 + 2, c2, subtitle, _font(9, False, MUTED), align="right").alignment = Alignment(horizontal="right", vertical="center", indent=1, readingOrder=2)
    ws.row_dimensions[r1 + 1].height = 24
    ws.row_dimensions[r1 + 2].height = 16


def _kpi(ws, col, kpi):
    c1, c2 = col, col + 3
    _box(ws, 5, c1, 8, c2)
    for c in range(c1, c2 + 1):
        ws.cell(5, c).fill = _fill(GOLD if kpi["key"] == "open" else TEAL)
    _merge(ws, 6, c1, 6, c2, "%s\n%s" % (kpi["ar"], kpi["en"]), _font(10, False, MUTED), align="right").alignment = Alignment(
        horizontal="right", vertical="center", wrap_text=True, indent=1, readingOrder=2)
    _merge(ws, 7, c1, 7, c2 - 1, kpi["value"], _font(26, True), align="right", fmt="#,##0").alignment = Alignment(
        horizontal="right", vertical="center", indent=1, readingOrder=2)
    g = kpi["growth"]
    if g is None:
        text, color, bg = "— عن الشهر السابق", MUTED, LINE
    else:
        up = g >= 0
        good = (up == kpi["good_up"])
        text = "%s %.1f%% عن الشهر السابق" % ("▲" if up else "▼", abs(g))
        color, bg = (GOOD, GOOD_BG) if (good or g == 0) else (BAD, BAD_BG)
    _merge(ws, 8, c1, 8, c2, text, _font(10, True, color), align="right", fill=bg).alignment = Alignment(
        horizontal="right", vertical="center", indent=1, readingOrder=2)
    ws.cell(8, c1).border = Border(top=Side(style="thin", color=LINE))


def _mini_line(ws_data, row, n):
    ch = LineChart()
    ch.add_data(Reference(ws_data, min_col=2, max_col=2, min_row=row, max_row=row + n - 1), titles_from_data=False)
    ch.legend = None
    ch.title = None
    ch.x_axis.delete = True
    ch.y_axis.delete = True
    ch.y_axis.majorGridlines = None
    ch.graphical_properties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    ch.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    s = ch.series[0]
    s.smooth = True
    s.graphicalProperties.line.solidFill = TEAL
    s.graphicalProperties.line.width = 22000
    s.marker.symbol = "none"
    ch.width, ch.height = 2.5, 1.9
    return ch


def build_workbook(data, rows, filters_text="كل البيانات"):
    wb = Workbook()
    ws = wb.active
    ws.title = "لوحة المؤشرات"
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL
    for c in range(1, COLS + 1):
        ws.column_dimensions[get_column_letter(c)].width = COL_W
    for r in range(1, 70):
        for c in range(1, COLS + 1):
            ws.cell(r, c).fill = _fill(PAGE_BG)

    # ---------------------------------------------------------------- data sheet (feeds the charts)
    wd = wb.create_sheet("البيانات")
    wd.sheet_view.rightToLeft = True
    wd.sheet_properties.tabColor = "8A8F98"
    head_font, head_fill = _font(10, True, "FFFFFF"), _fill(NAVY)

    def table(r0, c0, headers, body, widths=None):
        for j, h in enumerate(headers):
            cell = wd.cell(r0, c0 + j, h)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center", readingOrder=2)
        for i, row in enumerate(body, 1):
            for j, v in enumerate(row):
                cell = wd.cell(r0 + i, c0 + j, v)
                cell.font = _font(10)
                cell.alignment = Alignment(horizontal="right" if j == 0 else "center", readingOrder=2)
                if isinstance(v, (int, float)):
                    cell.number_format = "#,##0.##"
        return r0 + 1                                            # first data row

    months, bars = data["months"], data["bars"]
    r_months = table(1, 1, ["الشهر", "عدد الأوامر"], [("%s %s" % (m["label"], m["sub"]), b["count"]) for m, b in zip(months, bars)])
    sup = data["suppliers"]
    r_sup = table(10, 1, ["المورد", "إجمالي المصروف", "النسبة %"], [(s["name"], s["value"], s["share"]) for s in sup])
    proj = data["projects"]
    r_proj = table(19, 1, ["المشروع", "المصروف", "النسبة %"], [(p["name"], p["value"], p["share"]) for p in proj])
    stack = data["pstack"]
    status_order = data["status_order"]
    r_stack = table(30, 1, ["المشروع"] + status_order,
                    [[p["name"]] + [next((x["n"] for x in p["parts"] if x["status"] == st), 0) for st in status_order] for p in stack])
    trend = data["trend"]
    r_trend = table(42, 1, ["الشهر"] + [t["name"] for t in trend],
                    [["%s" % m["label"]] + [t["values"][i] for t in trend] for i, m in enumerate(months)])
    kp_row = 52
    table(kp_row, 1, ["المؤشر", "آخر 6 أشهر ←"] + [m["label"] for m in months], [])
    for i, k in enumerate(data["kpis"], 1):
        wd.cell(kp_row + i, 1, k["ar"]).font = _font(10)
        for j, v in enumerate(k["spark"]):
            wd.cell(kp_row + i, 3 + j, v).font = _font(10)
    for col, w in zip("ABCDEFGH", (34, 18, 18, 18, 18, 14, 14, 14)):
        wd.column_dimensions[col].width = w

    # ---------------------------------------------------------------- banner
    ws.row_dimensions[1].height = 8
    _merge(ws, 2, 1, 2, COLS, "لوحة مؤشرات أوامر الشراء", _font(22, True, "FFFFFF"), align="right", fill=NAVY).alignment = Alignment(
        horizontal="right", vertical="center", indent=2, readingOrder=2)
    _merge(ws, 3, 1, 3, COLS, "شركة الهضب للتجارة والمقاولات  ·  Purchase Orders Analytics  ·  %s  ·  صدر في %s" % (
        filters_text, datetime.now().strftime("%Y-%m-%d %H:%M")), _font(10, False, "CBD5E1"), align="right", fill=NAVY).alignment = Alignment(
        horizontal="right", vertical="center", indent=2, readingOrder=2)
    ws.row_dimensions[2].height = 40
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 10

    # ---------------------------------------------------------------- KPI cards (first card on the right, like the screen)
    ws.row_dimensions[5].height = 5
    ws.row_dimensions[6].height = 34
    ws.row_dimensions[7].height = 44
    ws.row_dimensions[8].height = 24
    for i, kpi in enumerate(data["kpis"]):
        _kpi(ws, 1 + i * 4, kpi)
        mini = _mini_line(wd, kp_row + 1 + i, len(kpi["spark"]))
        # the series lives in a ROW of the data sheet (one KPI per row)
        mini.series.clear()
        from openpyxl.chart import Series
        s = Series(Reference(wd, min_col=3, max_col=2 + len(kpi["spark"]), min_row=kp_row + 1 + i, max_row=kp_row + 1 + i), title=None)
        s.smooth = True
        s.graphicalProperties.line.solidFill = GOLD if kpi["key"] == "open" else TEAL
        s.graphicalProperties.line.width = 22000
        s.marker.symbol = "none"
        mini.series.append(s)
        ws.add_chart(mini, "%s7" % get_column_letter(4 + i * 4))
    ws.row_dimensions[9].height = 12

    # ---------------------------------------------------------------- panels
    A_END, B_START = 7, 8                                         # right block A-G (7 columns), left block H-P (9 columns)
    for r in range(10, 66):
        ws.row_dimensions[r].height = 20

    def anchor(col, row):
        return "%s%d" % (get_column_letter(col), row)

    # 1. orders growth - columns (newest bar dark, others sage) - right block
    _panel(ws, 10, 1, 27, A_END, "تطور عدد أوامر الشراء", "Purchase Orders Growth Trend · إجمالي الأوامر: %d" % data["total_count"])
    ch = _style_chart(BarChart())
    ch.type, ch.gapWidth = "col", 60
    ch.add_data(Reference(wd, min_col=2, min_row=1, max_row=1 + len(bars)), titles_from_data=True)
    ch.set_categories(Reference(wd, min_col=1, min_row=2, max_row=1 + len(bars)))
    s = ch.series[0]
    s.graphicalProperties.solidFill = SAGE
    s.graphicalProperties.line.noFill = True
    pt = DataPoint(idx=len(bars) - 1)
    pt.graphicalProperties.solidFill = TEAL
    s.dPt.append(pt)
    _labels(s, "0", "outEnd")
    _axes(ch, "0", grid=False)
    ch.y_axis.delete = True
    ch.x_axis.scaling.orientation = "maxMin"                    # oldest on the right, newest on the left (reads like the page)
    ch.width, ch.height = 15.3, 7.6
    ws.add_chart(ch, anchor(1, 13))

    # 2. top suppliers - horizontal bars - left block
    _panel(ws, 10, B_START, 27, COLS, "أكبر الموردين من حيث قيمة الصرف", "Top Suppliers by Expenditure")
    if sup:
        ch = _style_chart(BarChart())
        ch.type, ch.gapWidth = "bar", 70
        ch.add_data(Reference(wd, min_col=2, min_row=10, max_row=10 + len(sup)), titles_from_data=True)
        ch.set_categories(Reference(wd, min_col=1, min_row=11, max_row=10 + len(sup)))
        s = ch.series[0]
        s.graphicalProperties.solidFill = TEAL
        _labels(s, "#,##0", "outEnd")
        _axes(ch, grid=False)
        ch.y_axis.delete = True
        ch.x_axis.scaling.orientation = "maxMin"
        ch.width, ch.height = 19.6, 7.6
        ws.add_chart(ch, anchor(B_START, 13))
    else:
        _merge(ws, 17, B_START, 20, COLS, "لا توجد أوامر مصروفة بعد ضمن هذا الفلتر", _font(11, False, MUTED))

    # 3. project spending - doughnut - right block
    _panel(ws, 29, 1, 47, A_END, "توزيع المصروفات حسب المشروع", "Project Spending Distribution · الإجمالي %s ريال" % "{:,.0f}".format(data["paid_value"]))
    if proj:
        ch = DoughnutChart()
        _style_chart(ch)
        ch.holeSize = 58
        ch.legend = ch.legend or None
        from openpyxl.chart.legend import Legend
        ch.legend = Legend()
        ch.legend.position = "b"
        ch.add_data(Reference(wd, min_col=2, min_row=19, max_row=19 + len(proj)), titles_from_data=True)
        ch.set_categories(Reference(wd, min_col=1, min_row=20, max_row=19 + len(proj)))
        s = ch.series[0]
        for i in range(len(proj)):
            p = DataPoint(idx=i)
            p.graphicalProperties.solidFill = PALETTE[i % len(PALETTE)]
            p.graphicalProperties.line.solidFill = WHITE
            s.dPt.append(p)
        s.dLbls = DataLabelList()
        s.dLbls.showPercent = True
        for flag in ("showVal", "showSerName", "showCatName", "showLegendKey"):
            setattr(s.dLbls, flag, False)
        ch.width, ch.height = 15.3, 8.6
        ws.add_chart(ch, anchor(1, 32))
    else:
        _merge(ws, 36, 1, 40, A_END, "لا توجد مصروفات مسجّلة بعد", _font(11, False, MUTED))

    # 4. supplier payment trends - left block
    _panel(ws, 29, B_START, 47, COLS, "تحليل الإنفاق الشهري للموردين", "Supplier Payment Trends")
    if trend:
        ch = _style_chart(AreaChart())
        ch.grouping = "standard"
        ch.add_data(Reference(wd, min_col=2, max_col=1 + len(trend), min_row=42, max_row=42 + len(months)), titles_from_data=True)
        ch.set_categories(Reference(wd, min_col=1, min_row=43, max_row=42 + len(months)))
        colors = [TEAL, GOLD, "6F9EA3", "8A8F98"]
        for i, s in enumerate(ch.series):
            s.graphicalProperties.solidFill = colors[i % 4]
            s.graphicalProperties.line.solidFill = colors[i % 4]
            s.graphicalProperties.line.width = 22000
        from openpyxl.chart.legend import Legend
        ch.legend = Legend()
        ch.legend.position = "b"
        _axes(ch)
        ch.x_axis.scaling.orientation = "maxMin"
        ch.width, ch.height = 19.6, 8.6
        ws.add_chart(ch, anchor(B_START, 32))
    else:
        _merge(ws, 36, B_START, 40, COLS, "لا توجد مدفوعات لعرض اتجاهها بعد", _font(11, False, MUTED))

    # 5. orders by project, stacked by status - full width
    _panel(ws, 49, 1, 65, COLS, "توزيع أوامر الشراء حسب المشروع", "Purchase Orders by Project · مكتمل / بانتظار السداد / بانتظار الفاتورة / بانتظار الاستلام")
    if stack:
        ch = _style_chart(BarChart())
        ch.type, ch.grouping, ch.overlap, ch.gapWidth = "bar", "stacked", 100, 50
        ch.add_data(Reference(wd, min_col=2, max_col=1 + len(status_order), min_row=30, max_row=30 + len(stack)), titles_from_data=True)
        ch.set_categories(Reference(wd, min_col=1, min_row=31, max_row=30 + len(stack)))
        for s, st in zip(ch.series, status_order):
            s.graphicalProperties.solidFill = STATUS_COLORS[st]
            s.graphicalProperties.line.noFill = True
        from openpyxl.chart.legend import Legend
        ch.legend = Legend()
        ch.legend.position = "t"
        _axes(ch, "0")
        ch.x_axis.scaling.orientation = "maxMin"
        ch.width, ch.height = 35.2, 7.4
        ws.add_chart(ch, anchor(1, 52))
    else:
        _merge(ws, 55, 1, 60, COLS, "لا توجد أوامر شراء ضمن هذا الفلتر", _font(11, False, MUTED))

    _merge(ws, 67, 1, 67, COLS, "أُنشئ هذا التقرير من نظام إدارة العقود والمشتريات — شركة الهضب للتجارة والمقاولات", _font(9, False, MUTED), align="center")

    # ---------------------------------------------------------------- print setup: one A4 landscape page
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:%s67" % get_column_letter(COLS)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.35
    ws.oddFooter.center.text = "شركة الهضب للتجارة والمقاولات — صفحة &P من &N"

    # ---------------------------------------------------------------- the orders themselves
    wo = wb.create_sheet("أوامر الشراء")
    wo.sheet_view.rightToLeft = True
    wo.sheet_properties.tabColor = SAGE
    heads = ["رقم الأمر", "التاريخ", "المورد", "المشروع", "الإجمالي", "حالة الاستلام", "حالة الفاتورة", "حالة السداد", "تاريخ السداد"]
    for j, h in enumerate(heads, 1):
        c = wo.cell(1, j, h)
        c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center", readingOrder=2)
    for i, r in enumerate(sorted(rows, key=lambda r: r.get("OrderDate") or datetime.min, reverse=True), 2):
        vals = [r.get("PONumber"), r["OrderDate"].strftime("%Y-%m-%d") if r.get("OrderDate") else "", r.get("SupplierName"), r.get("Project"),
                r.get("TotalAmount") or 0, r.get("ReceiptStatus"), r.get("InvoiceStatus"), r.get("PaymentStatus"),
                r["PaymentDate"].strftime("%Y-%m-%d") if r.get("PaymentDate") else ""]
        for j, v in enumerate(vals, 1):
            c = wo.cell(i, j, v)
            c.font = _font(10)
            c.alignment = Alignment(horizontal="right" if j in (3, 4) else "center", readingOrder=2)
            if j == 5:
                c.number_format = "#,##0.00"
            if i % 2 == 1:
                c.fill = _fill("F8FAFC")
    for j, w in enumerate((16, 13, 34, 26, 16, 16, 16, 16, 14), 1):
        wo.column_dimensions[get_column_letter(j)].width = w
    wo.freeze_panes = "A2"
    wo.auto_filter.ref = "A1:I%d" % max(2, len(rows) + 1)
    wo.page_setup.orientation = "landscape"
    wo.page_setup.fitToWidth, wo.page_setup.fitToHeight = 1, 0
    wo.sheet_properties.pageSetUpPr.fitToPage = True
    return wb
