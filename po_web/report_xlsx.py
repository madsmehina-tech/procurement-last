# -*- coding: utf-8 -*-
"""Professional Excel workbooks for the management reports: a title band, then each sheet is a real Excel TABLE (Pivot / Power BI read them by name)."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def build(sheets, label, company="شركة الهضب للتجارة والمقاولات"):
    """sheets: [{title, band, headers, rows, widths, money (1-based columns), table}] -> BytesIO."""
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(sh["title"][:31])
        ws.sheet_view.rightToLeft = True
        ws.sheet_view.showGridLines = False
        n = len(sh["headers"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
        ws["A1"] = "%s — %s" % (company, sh["band"])
        ws["A1"].font, ws["A1"].fill = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF"), PatternFill("solid", start_color="14323A")
        ws["A1"].alignment = Alignment(horizontal="right", vertical="center", indent=1, readingOrder=2)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
        ws["A2"] = "%s  ·  صدر في %s" % (label or "كل البيانات", datetime.now().strftime("%Y-%m-%d %H:%M"))
        ws["A2"].font, ws["A2"].fill = Font(name="Segoe UI", size=10, color="CBD5E1"), PatternFill("solid", start_color="14323A")
        ws["A2"].alignment = Alignment(horizontal="right", vertical="center", indent=1, readingOrder=2)
        ws.row_dimensions[1].height, ws.row_dimensions[2].height = 28, 18
        for j, h in enumerate(sh["headers"], 1):
            ws.cell(4, j, h)
        money = set(sh.get("money", ()))
        for i, row in enumerate(sh["rows"], 5):
            for j, v in enumerate(row, 1):
                c = ws.cell(i, j, v)
                c.font = Font(name="Segoe UI", size=10)
                c.alignment = Alignment(horizontal="center" if (j in money or isinstance(v, (int, float))) else "right", vertical="center", readingOrder=2, wrap_text=isinstance(v, str) and len(v) > 40)
                if j in money:
                    c.number_format = "#,##0.00"
        last = 4 + max(len(sh["rows"]), 1)
        if sh["rows"]:
            t = Table(displayName=sh["table"], ref="A4:%s%d" % (get_column_letter(n), last))
            t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)
        else:
            for j in range(1, n + 1):
                ws.cell(4, j).font = Font(bold=True)
        for j, w in enumerate(sh["widths"], 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "A5"
        ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", 9
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "4:4"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
