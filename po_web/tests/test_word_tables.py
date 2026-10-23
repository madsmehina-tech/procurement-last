# -*- coding: utf-8 -*-
"""The tables the system builds into the contracts must follow the OOXML property order, or Word may call the file corrupt (audit M-26)."""
import conftest  # noqa: F401
import contract_render as r
from docx import Document


def _kids(el):
    return [c.tag.split("}")[-1] for c in el]


def _table():
    doc = Document()
    spec = {"header": ["البند", "السعر"], "rows": [["حفر", "100.00"], [None, "5.00"]], "widths": [3, 1], "align": ["right", "center"], "blank_missing": True, "bold_last": True}
    return doc, r.build_table(doc, spec, 9000)


def _in_order(kids, order):
    rank = {n: i for i, n in enumerate(order.split())}
    pos = [rank[k] for k in kids if k in rank]
    return pos == sorted(pos)


def test_table_properties_are_in_standard_order():
    _doc, t = _table()
    assert _in_order(_kids(t._tbl.tblPr), r._ORDER["tblPr"]), _kids(t._tbl.tblPr)
    for row in t.rows:
        for cell in row.cells:
            assert _in_order(_kids(cell._tc.tcPr), r._ORDER["tcPr"]), _kids(cell._tc.tcPr)
            for mar in cell._tc.tcPr.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcMar"):
                assert _in_order(_kids(mar), r._ORDER["tcMar"])
            for run in cell.paragraphs[0].runs:
                assert _in_order(_kids(run._r.rPr), r._ORDER["rPr"]), _kids(run._r.rPr)
            assert _in_order(_kids(cell.paragraphs[0]._p.pPr), r._ORDER["pPr"])


def test_grid_columns_carry_the_real_widths():
    _doc, t = _table()
    grid = [int(g.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w")) for g in t._tbl.tblGrid]
    assert grid[0] > grid[1] * 2.5                                       # 3 : 1, not equal columns


def test_the_document_still_opens_and_keeps_its_content(tmp_path):
    doc, t = _table()
    out = tmp_path / "t.docx"
    doc.save(out)
    again = Document(str(out))
    assert again.tables[0].rows[1].cells[0].text == "حفر"


def test_every_shipped_starter_template_renders_without_leftover_marks(tmp_path):
    import json
    import os
    from docx import Document as D
    starters = os.path.join(os.path.dirname(os.path.abspath(r.__file__)), "contract_starters")
    for s in json.load(open(os.path.join(starters, "starters.json"), encoding="utf-8")):
        out = tmp_path / (s["code"] + ".docx")
        tables = {n: {"header": ["أ", "ب"], "rows": [["1", "2"]], "widths": [1, 1]} for n in ("جدول_الأسعار", "جدول_الدفعات", "جدول_المواصفات")}
        res = r.render(os.path.join(starters, s["file"]), str(out), {}, tables)
        text = "\n".join(p.text for p in D(str(out)).paragraphs)
        assert "{{" not in text and "[[" not in text, s["code"]
        assert res["missing"] is not None
