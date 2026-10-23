# -*- coding: utf-8 -*-
"""
contract_render.py - fills a Word template (.docx): the {{ }} marks become values, the [[ ]] marks become tables.
A value that is missing stays as dots (........) HIGHLIGHTED YELLOW so the person sees what to complete.
The template's own look (letterhead in the header / footer, fonts, spacing) is kept: only the marks change.
"""
import copy

import contract_vars as V

try:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Twips
    DOCX_ERROR = None
except ImportError as _e:                               # the system still starts; only the contracts pages say what is missing
    Document = WD_COLOR_INDEX = OxmlElement = qn = Pt = Twips = None
    DOCX_ERROR = _e


class MissingDocx(RuntimeError):
    """python-docx is not installed on THIS computer (libraries are installed per computer)."""


def need_docx():
    if DOCX_ERROR is not None:
        raise MissingDocx("مكتبة python-docx غير مثبّتة على هذا الجهاز. افتح cmd واكتب:  cd /d E:\\PO_System\\po_web  ثم  python -m pip install -r requirements.txt  وأعد تشغيل النظام.")


# ---------------------------------------------------------------- walking the document
def _tables_paragraphs(tables):
    for t in tables:
        for row in t.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                yield from cell.paragraphs
                yield from _tables_paragraphs(cell.tables)


def iter_paragraphs(doc):
    yield from doc.paragraphs
    yield from _tables_paragraphs(doc.tables)
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer, s.even_page_header, s.even_page_footer):
            try:
                if part is not None and not part.is_linked_to_previous:
                    yield from part.paragraphs
                    yield from _tables_paragraphs(part.tables)
            except Exception:                           # noqa - a header that cannot be read is skipped
                continue


# ---------------------------------------------------------------- replacing text across runs
def _split_at(p, offset):
    """Makes `offset` (a character position in the paragraph's text) a run boundary."""
    pos = 0
    for r in p.runs:
        n = len(r.text)
        if pos < offset < pos + n:
            k = offset - pos
            right = copy.deepcopy(r._r)
            r._r.addnext(right)
            from docx.text.run import Run
            right_run = Run(right, p)
            text = r.text
            r.text, right_run.text = text[:k], text[k:]
            return
        pos += n


def _replace_span(p, start, end, text, highlight):
    _split_at(p, start)
    _split_at(p, end)
    pos, inside = 0, []
    for r in p.runs:
        n = len(r.text)
        if pos >= start and pos + n <= end and n > 0:
            inside.append(r)
        pos += n
    if not inside:
        return
    first = inside[0]
    first.text = text
    for r in inside[1:]:
        r._r.getparent().remove(r._r)
    if highlight:
        first.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _remove_paragraph(p):
    """Removes a paragraph (a table cell must keep one: it is emptied instead)."""
    parent = p._p.getparent()
    if parent.tag == qn("w:tc") and len(parent.findall(qn("w:p"))) == 1:
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        return
    parent.remove(p._p)


def fill_paragraph(p, values, missing):
    """Replaces every {{ }} mark of the paragraph.
    * a LIST mark (قائمة_...) with several lines repeats the paragraph once per line;
    * an OPTIONAL mark (..._اختياري) that is empty removes the paragraph when it holds nothing else, else leaves nothing (never yellow);
    * any other empty mark stays as dots highlighted yellow and its key goes to `missing`."""
    matches = list(V.FIELD_RE.finditer(p.text))
    if not matches:
        return
    for m in matches:                                   # a list with several items: one paragraph per item
        key = V.norm_key(m.group(1))
        if V.is_list(key):
            lines = [x.strip() for x in str(values.get(key) or "").splitlines() if x.strip()]
            if len(lines) > 1:
                from docx.text.paragraph import Paragraph
                pars, prev = [p], p._p
                for _ in lines[1:]:
                    clone = copy.deepcopy(p._p)
                    prev.addnext(clone)
                    prev = clone
                    pars.append(Paragraph(clone, p._parent))
                for par, line in zip(pars, lines):
                    fill_paragraph(par, dict(values, **{key: line}), missing)
                return
    for m in reversed(matches):
        key = V.norm_key(m.group(1))
        val = values.get(key)
        empty = val is None or str(val).strip() == ""
        if empty and V.is_optional(key):
            rest = (p.text[:m.start()] + p.text[m.end():]).strip(" \t•-–:،.\u200f\u00a0")
            if not rest and len(matches) == 1:
                _remove_paragraph(p)
                return
            _replace_span(p, m.start(), m.end(), "", False)
            continue
        if empty:
            missing.add(key)
        _replace_span(p, m.start(), m.end(), V.BLANK if empty else str(val).strip(), empty)


# ---------------------------------------------------------------- tables built by the system
FONT = "Sakkal Majalla"                                # the company own contracts are written in it (Windows has it)
TABLE_SIZE = 22


def _rpr_rtl(run, size=20, bold=False):
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(a), FONT)
    for tag, val in (("w:sz", size), ("w:szCs", size)):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), str(val))
    if bold:
        for tag in ("w:b", "w:bCs"):
            if rpr.find(qn(tag)) is None:
                rpr.append(OxmlElement(tag))
    if rpr.find(qn("w:rtl")) is None:
        rpr.append(OxmlElement("w:rtl"))


def _cell(cell, text, width, align, head=False, blank=False, bold=False):
    cell.width = Twips(width)
    tcpr = cell._tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "center")
    tcpr.append(va)
    if head:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "E7E6E6")
        tcpr.append(shd)
    mar = OxmlElement("w:tcMar")
    for side, w in (("top", 60), ("bottom", 60), ("start", 100), ("end", 100)):
        e = OxmlElement("w:%s" % side)
        e.set(qn("w:w"), str(w))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcpr.append(mar)
    para = cell.paragraphs[0]
    ppr = para._p.get_or_add_pPr()
    ppr.append(OxmlElement("w:bidi"))
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), "0")
    sp.set(qn("w:after"), "0")
    ppr.append(sp)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), {"center": "center", "right": "left", "left": "right"}.get(align, "center"))   # in a bidi paragraph "left" is the visual right edge
    ppr.append(jc)
    missing = text is None and blank                    # None = a value that is missing (dots, yellow); "" = intentionally empty
    run = para.add_run(V.BLANK if missing else (text or ""))
    _rpr_rtl(run, TABLE_SIZE, head or bold)
    if missing:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def build_table(doc, spec, total_width):
    """spec: {header: [...], rows: [[...]], widths: [fractions], align: [center|right], blank_missing: bool, bold_last: bool}"""
    cols = len(spec["header"])
    rows = spec["rows"]
    t = doc.add_table(rows=1 + len(rows), cols=cols)
    tbl = t._tbl
    tblpr = tbl.tblPr
    for tag in ("w:tblStyle", "w:tblW", "w:tblLayout", "w:tblBorders", "w:bidiVisual", "w:jc"):
        for el in tblpr.findall(qn(tag)):
            tblpr.remove(el)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(total_width))
    w.set(qn("w:type"), "dxa")
    tblpr.append(w)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    tblpr.append(jc)
    tblpr.append(OxmlElement("w:bidiVisual"))
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement("w:%s" % side)
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "808080")
        borders.append(b)
    tblpr.append(borders)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    fr = spec.get("widths") or [1.0 / cols] * cols
    widths = [int(total_width * f / sum(fr)) for f in fr]
    aligns = spec.get("align") or ["center"] * cols
    for j, h in enumerate(spec["header"]):
        _cell(t.rows[0].cells[j], h, widths[j], "center", head=True)
    trpr = t.rows[0]._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:tblHeader"))
    for i, row in enumerate(rows, 1):
        last = spec.get("bold_last") and i == len(rows)
        for j, v in enumerate(row):
            _cell(t.rows[i].cells[j], v, widths[j], aligns[j], blank=spec.get("blank_missing", True), bold=last)
        cant = OxmlElement("w:cantSplit")
        t.rows[i]._tr.get_or_add_trPr().append(cant)
    return t


def _usable_width(doc):
    s = doc.sections[0]
    return int((s.page_width - s.left_margin - s.right_margin) / 635)        # EMU -> twips


def render(template_path, out_path, values, tables):
    """Fills the template. `tables` maps a block name to a spec (see build_table). Returns {"missing": [keys], "blocks_used": [...], "unknown_blocks": [...]}."""
    need_docx()
    doc = Document(template_path)
    missing = set()
    used, unknown = [], []
    width = _usable_width(doc)
    for p in list(doc.paragraphs):                      # the [[ ]] blocks (a whole paragraph) first
        m = V.BLOCK_RE.fullmatch(p.text.strip()) if p.text.strip() else None
        if not m:
            continue
        name = V.norm_key(m.group(1))
        spec = tables.get(name)
        if spec is None:
            unknown.append(name)
            continue
        t = build_table(doc, spec, width)
        p._p.addnext(t._tbl)
        p._p.getparent().remove(p._p)
        used.append(name)
    for p in list(iter_paragraphs(doc)):
        if "{{" in p.text:
            fill_paragraph(p, values, missing)
    doc.save(out_path)
    return {"missing": sorted(missing), "blocks_used": used, "unknown_blocks": unknown}
