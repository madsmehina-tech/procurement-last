# -*- coding: utf-8 -*-
"""Subcontractors: the list (serial, filters, Excel / PDF / print) and THE page that manages everything about one contractor:
its data, papers, contracts (add / view / edit / delete), certificates (add / edit / delete) and archive files."""
import io
import logging
from datetime import datetime

from flask import flash, redirect, render_template, request, send_file, url_for

import archive
import auth
import contracts_service as cs
import cycle_files
import db
import ledger_service
import pdfgen
from finance import STATUS_BADGE
from routes import route
from textnorm import clean_spaces, contains, norm_key

SLOTS_TOTAL = 4        # commercial register, tax certificate, IBAN certificate, qualification


def _papers_state():
    out = {}
    for d in db.query("SELECT [EntityKey], [DocType], [IsCurrent], [ExpiryDate] FROM [Documents] WHERE [EntityType] = ?", ("SUB",)):
        if not d["IsCurrent"]:
            continue
        st = out.setdefault(str(d["EntityKey"]), {"slots": set(), "worst": None})
        if d["DocType"] in archive.SUB_SLOTS:
            st["slots"].add(d["DocType"])
        state = archive.expiry_state(d.get("ExpiryDate"))
        if state and state[0] in ("expired", "soon") and (not st["worst"] or state[0] == "expired"):
            st["worst"] = state
    return out


def _form():
    f = request.form
    return (clean_spaces(f.get("name", "")), f.get("cr", "").strip(), f.get("tax_number", "").strip(), f.get("bank_account", "").strip(),
            f.get("phone", "").strip(), f.get("notes", "").strip())


# --------------------------------------------------------------------------- the list
def _list_rows():
    q = request.args.get("q", "").strip()
    subs, contracts, certs, _p = cs.load()
    papers = _papers_state()
    rows = []
    for sid, s in subs.items():
        if q and not (contains(s["Name"], q) or contains(s.get("CR") or "", q) or contains(s.get("TaxNumber") or "", q)):
            continue
        p = papers.get(str(sid), {"slots": set(), "worst": None})
        rows.append(dict(s, **cs.sub_figures(sid, contracts, certs), files=len(p["slots"]), worst=p["worst"]))
    rows.sort(key=lambda r: norm_key(r["Name"]))
    for i, r in enumerate(rows, 1):
        r["serial"] = i
    kpi = {"count": len(rows), "due": sum(r["due"] for r in rows), "paid": sum(r["paid"] for r in rows), "remaining": sum(r["remaining"] for r in rows)}
    return rows, kpi, q


@route("/subcontractors")
@auth.login_required
def subcontractors_list():
    rows, kpi, q = _list_rows()
    return render_template("subcontractors_list.html", rows=rows, kpi=kpi, q=q, slots_total=SLOTS_TOTAL)


@route("/subcontractors/export.xlsx")
@auth.login_required
def subcontractors_export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    rows, kpi, q = _list_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "مقاولو الباطن"
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    heads = ["مسلسل", "المقاول", "السجل التجاري", "الرقم الضريبي", "جوال المسؤول", "العقود", "المستخلصات", "قيمة العقود", "المستحق", "المصروف", "المتبقي", "الأوراق"]
    navy, line = PatternFill("solid", start_color="1E293B"), Side(style="thin", color="E2E8F0")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(heads))
    ws["A1"] = "سجل مقاولي الباطن — شركة الهضب للتجارة والمقاولات"
    ws["A1"].font, ws["A1"].fill = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF"), navy
    ws["A1"].alignment = Alignment(horizontal="right", vertical="center", indent=1, readingOrder=2)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(heads))
    ws["A2"] = "%s  ·  %d مقاول  ·  المستحق %s  ·  المصروف %s  ·  المتبقي %s  ·  صدر في %s" % (("بحث: " + q) if q else "كل المقاولين", kpi["count"], "{:,.0f}".format(kpi["due"]),
                                                                                                "{:,.0f}".format(kpi["paid"]), "{:,.0f}".format(kpi["remaining"]), datetime.now().strftime("%Y-%m-%d %H:%M"))
    ws["A2"].font, ws["A2"].fill = Font(name="Segoe UI", size=10, color="CBD5E1"), navy
    ws["A2"].alignment = Alignment(horizontal="right", vertical="center", indent=1, readingOrder=2)
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 30, 20
    for j, h in enumerate(heads, 1):
        c = ws.cell(4, j, h)
        c.font, c.fill = Font(name="Segoe UI", bold=True, color="FFFFFF"), PatternFill("solid", start_color="29595E")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
    for i, r in enumerate(rows, 1):
        vals = [r["serial"], r["Name"], r.get("CR") or "", r.get("TaxNumber") or "", r.get("Phone") or "", r["contracts"], r["certs"], r["value"], r["due"], r["paid"], r["remaining"],
                "%d / %d" % (r["files"], SLOTS_TOTAL)]
        for j, v in enumerate(vals, 1):
            c = ws.cell(4 + i, j, v)
            c.font, c.border = Font(name="Segoe UI", size=10), Border(bottom=line)
            c.alignment = Alignment(horizontal="right" if j == 2 else "center", readingOrder=2)
            if j in (8, 9, 10, 11):
                c.number_format = "#,##0.00"
            if i % 2 == 0:
                c.fill = PatternFill("solid", start_color="F8FAFC")
    last = 4 + len(rows)
    ws.cell(last + 1, 2, "الإجمالي").font = Font(name="Segoe UI", bold=True)
    for j, key in ((8, "value"), (9, "due"), (10, "paid"), (11, "remaining")):
        c = ws.cell(last + 1, j, round(sum(r[key] for r in rows), 2))
        c.font, c.number_format = Font(name="Segoe UI", bold=True), "#,##0.00"
    for j, w in enumerate((8, 34, 16, 16, 15, 9, 11, 16, 16, 16, 16, 10), 1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = "A5"
    ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", 9
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="subcontractors_%s.xlsx" % datetime.now().strftime("%Y%m%d"),
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _report_html(pdf=False):
    rows, kpi, q = _list_rows()
    return render_template("subcontractors_report.html", rows=rows, kpi=kpi, label=("بحث: " + q) if q else "كل المقاولين",
                           now=datetime.now().strftime("%Y-%m-%d %H:%M"), logo_src=cycle_files.logo() if pdf else "", slots_total=SLOTS_TOTAL)


@route("/subcontractors/print")
@auth.login_required
def subcontractors_print():
    return _report_html()


@route("/subcontractors/export.pdf")
@auth.login_required
def subcontractors_export_pdf():
    html = _report_html(pdf=True)
    if pdfgen.available():
        import os
        import tempfile
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        try:
            pdfgen.html_to_pdf(html, path)
            with open(path, "rb") as fh:
                data = fh.read()
            return send_file(io.BytesIO(data), as_attachment=True, download_name="subcontractors_%s.pdf" % datetime.now().strftime("%Y%m%d"), mimetype="application/pdf")
        except pdfgen.PdfError as e:
            flash("تعذّر إنشاء PDF تلقائيًا (%s) - استخدم «طباعة» ثم «حفظ كـ PDF» من المتصفح." % e)
        finally:
            try:
                os.remove(path)
            except OSError:
                pass
    return html


# --------------------------------------------------------------------------- create
@route("/subcontractors/new", methods=["GET", "POST"])
@auth.login_required
def subcontractor_new():
    if request.method == "GET":
        return render_template("subcontractor_form.html", sub=None)
    name, cr, tax, bank, phone, notes = _form()
    if not name:
        flash("اسم المقاول مطلوب.")
        return render_template("subcontractor_form.html", sub=request.form)
    if any(norm_key(r["Name"]) == norm_key(name) for r in db.query("SELECT [Name] FROM [Subcontractors]")):
        flash("يوجد مقاول بهذا الاسم (أو اسم مطابق له). افتح ملفه بدل تكراره.")
        return render_template("subcontractor_form.html", sub=request.form)
    try:
        sub_id = db.execute_insert(
            "INSERT INTO [Subcontractors] ([Name], [CR], [TaxNumber], [BankAccount], [Phone], [Notes], [CreatedAt]) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (name, cr, tax, bank, phone, notes, datetime.now()))
    except Exception as e:                                 # noqa
        logging.getLogger("po.error").exception("subcontractor save failed")
        flash("تعذّر حفظ المقاول ولم يتغير أي شيء. السبب: %s" % " ".join(str(e).split())[:200])
        return render_template("subcontractor_form.html", sub=request.form)
    archive.ensure_sub_folder(sub_id)                     # إدارة مقاولي الباطن / <المقاول> / {أوراق المقاول، العقود، المستخلصات}
    flash("تم إضافة المقاول. من هذه الصفحة تُدار أوراقه وعقوده ومستخلصاته.")
    return redirect(url_for("subcontractor_detail", sub_id=sub_id))


# --------------------------------------------------------------------------- the contractor's page
def _next_actions(sub_id, my_certs, contracts_rows, ledger, url_for_):
    """The «what to do next» strip at the top of the contractor page: [{level, text, tab}] most urgent first."""
    out = []
    entered = [g for g in my_certs if g["Status"] == "مُدخَل" and g["Remaining"] > 0.005 and not g.get("ClosedAt")]
    if entered:
        out.append({"level": "warning", "tab": "certificates", "text": "%d مستخلص مُدخَل بقيمة %s ريال لم يدخل دورة صرف" % (len(entered), "{:,.0f}".format(sum(g["Remaining"] for g in entered)))})
    in_cycle = [g for g in my_certs if g["Status"] == "ضمن دورة صرف" and not g.get("ClosedAt")]
    if in_cycle:
        out.append({"level": "info", "tab": "certificates", "text": "%d مستخلص ضمن دورة صرف حالية" % len(in_cycle)})
    waiting = [e for e in ledger if e.get("Kind") == ledger_service.K_CREDIT and e.get("Category") == ledger_service.CAT_PAYMENT and e.get("Status") != ledger_service.S_DONE]
    if waiting:
        out.append({"level": "warning", "tab": "ledger", "text": "%d دفعة مباشرة مسجّلة ولم تُنفَّذ بعد" % len(waiting)})
    if not any((c.get("Status") or "ساري") == "ساري" for c in contracts_rows):
        out.append({"level": "info", "tab": "contracts", "text": "لا يوجد عقد ساري لهذا المقاول"})
    order = {"danger": 0, "warning": 1, "info": 2}
    return sorted(out, key=lambda a: order[a["level"]])


def _summary(sub_id, contracts_rows, my_certs, pays, fig):
    """The financial summary shown at the top: what management asks first about any contractor."""
    live = [g for g in my_certs if g["Status"] != "مرفوض"]
    my_pays = [p for p in pays if p.get("SubcontractorID") == sub_id or p["CertificateID"] in {g["CertificateID"] for g in my_certs}]
    last_pay = max(my_pays, key=lambda p: (p.get("PaidAt") or datetime.min, p["PaymentID"])) if my_pays else None
    direct = [e for e in ledger_service.entries(sub_id) if e["Kind"] == ledger_service.K_CREDIT and e["Category"] == ledger_service.CAT_PAYMENT and e["Status"] == ledger_service.S_DONE]
    last_direct = max(direct, key=lambda e: (e.get("EntryDate") or datetime.min, e["EntryID"])) if direct else None
    by_cert = {g["CertificateID"]: g for g in my_certs}
    last_cert = max(my_certs, key=lambda g: (g.get("CertificateDate") or g.get("CreatedAt") or datetime.min, g["CertificateID"])) if my_certs else None
    open_ = [g for g in live if g["Status"] in ("مُدخَل", "ضمن دورة صرف") and g["Remaining"] > 0.005 and not g.get("ClosedAt")]
    return {"open_total": round(sum(g["Remaining"] for g in open_), 2), "open_count": len(open_),
            "contracts_value": round(sum(c["value"] for c in contracts_rows), 2), "certs_value": round(sum(g["Value"] for g in live), 2),
            "paid": fig["paid"], "remaining": fig["remaining"],     # the SAME figures as the account statement (see contracts_service.sub_figures)
            "n_contracts": len(contracts_rows), "n_certs": len(my_certs),
            "last_pay": (last_direct and {"amount": last_direct["Amount"] or 0, "date": last_direct.get("EntryDate"), "cert": None})
            if (last_direct and (not last_pay or (last_direct.get("EntryDate") or datetime.min) >= (last_pay.get("PaidAt") or datetime.min)))
            else (last_pay and {"amount": last_pay.get("Amount") or 0, "date": last_pay.get("PaidAt"), "cert": by_cert.get(last_pay["CertificateID"])}),
            "last_cert": last_cert}


def project_site_lists():
    """The unified lists every certificate and contract form offers: ([project names], {project: [site names]}).
    A project or site typed for the first time is added to them when the record is saved."""
    projects = {p["ProjectID"]: p["ProjectName"] for p in db.query("SELECT [ProjectID], [ProjectName], [IsActive] FROM [Projects]") if p.get("IsActive") is not False}
    by_project = {name: [] for name in projects.values()}
    for st in db.query("SELECT [ProjectID], [SiteName], [IsActive] FROM [Sites]"):
        if st.get("IsActive") is not False and st["ProjectID"] in projects:
            by_project[projects[st["ProjectID"]]].append(st["SiteName"])
    return sorted(by_project, key=norm_key), {k: sorted(v, key=norm_key) for k, v in by_project.items()}


def render_sub_page(sub_id, open_dialog=None, form=None):
    subs, contracts, certs, pays = cs.load()
    sub = subs.get(sub_id)
    if not sub:
        flash("لم يتم العثور على المقاول.")
        return redirect(url_for("subcontractors_list"))
    mine = sorted((c for c in contracts.values() if c.get("SubcontractorID") == sub_id), key=lambda c: c.get("CreatedAt") or datetime.min, reverse=True)
    # ONE query per kind (not one per contract / certificate): on Access each query costs a new connection, and a contractor can have hundreds of certificates
    newest_first = lambda d: ((d.get("UploadedAt") or datetime.min), d["DocID"])
    docs_by = {"CONTRACT": {}, "CERT": {}}
    for kind in docs_by:
        for d in archive.list_documents_many(kind):
            docs_by[kind].setdefault(str(d["EntityKey"]), []).append(d)
    sites_of = {}
    for s_ in db.query("SELECT [ContractID], [SiteName] FROM [ContractSites]"):
        sites_of.setdefault(s_["ContractID"], []).append(s_["SiteName"])
    rows = []
    for c in mine:
        sites = list(sites_of.get(c["ContractID"], []))
        if c.get("SiteName") and c["SiteName"] not in sites:
            sites.insert(0, c["SiteName"])
        files = sorted(docs_by["CONTRACT"].get(str(c["ContractID"]), []), key=newest_first, reverse=True)
        rows.append(dict(c, sites=sites, files=files, **cs.contract_figures(c, certs)))
    my_certs = sorted((g for g in certs if g["SubcontractorID"] == sub_id), key=lambda g: (g.get("CertificateDate") or g.get("CreatedAt") or datetime.min, g["CertificateID"]), reverse=True)
    cert_files = {}
    for g in my_certs:
        cert_files[g["CertificateID"]] = {d["DocType"]: d for d in sorted(docs_by["CERT"].get(str(g["CertificateID"]), []), key=newest_first, reverse=True)}
    projects_list, sites_by_project = project_site_lists()
    contract_choices = [{"id": c["ContractID"], "label": "%s — %s" % (c.get("ContractNo") or ("#%s" % c["ContractID"]), c.get("Project") or "")} for c in mine]
    next_numbers = {}                                # "<project>|<site>" -> the next certificate number of THIS contractor there
    for g in my_certs:
        key = "%s|%s" % (g.get("Project") or "", g.get("SiteName") or "")
        next_numbers[key] = max(next_numbers.get(key, 0), g["CertificateNo"] or 0)
    next_numbers = {k: v + 1 for k, v in next_numbers.items()}
    slots = {sl["type"]: sl for sl in archive.panel("SUB", sub_id, archive.SUB_DOC_TYPES)}
    others = sorted((d for t in archive.SUB_OTHER_TYPES for d in slots[t]["current"]), key=lambda d: d["UploadedAt"] or datetime.min, reverse=True)
    filled = sum(1 for t in archive.SUB_SLOTS if slots[t]["current"])
    folder = (sub.get("ArchiveFolder") or archive.sub_root_rel(sub["Name"])) + "/" + archive.SUB_PAPERS
    ledger_files = {}
    for d in archive.list_documents_many("LEDGER") if hasattr(archive, "list_documents_many") else []:
        ledger_files.setdefault(int(d["EntityKey"]), []).append(d)
    fig = cs.sub_figures(sub_id, contracts, certs)
    ledger_rows = ledger_service.entries(sub_id)
    return render_template("subcontractor_detail.html", sub=sub, fig=fig, next_actions=_next_actions(sub_id, my_certs, rows, ledger_rows, url_for), summary=_summary(sub_id, rows, my_certs, pays, fig),
                           contracts=rows, certs=my_certs, cert_files=cert_files, contract_choices=contract_choices, sites_by_project=sites_by_project, next_numbers=next_numbers, ledger=ledger_rows, ledger_files=ledger_files, balance=ledger_service.balance_of(sub_id), ledger_categories=ledger_service.CATEGORIES, badge=STATUS_BADGE, slots=slots, others=others, filled=filled,
                           slots_total=SLOTS_TOTAL, slot_names=archive.SUB_SLOTS, other_types=archive.SUB_OTHER_TYPES, folder=folder,
                           network=archive.network_path(folder), statuses=["ساري", "موقوف", "منتهي"], projects_list=projects_list,
                           open_dialog=open_dialog or _open_from_query(), form=form or {}, today_iso=datetime.now().strftime("%Y-%m-%d"),
                           root_folder=archive.sub_root_rel(sub["Name"]))


def _open_from_query():
    a = request.args
    if a.get("new_contract"):
        return "contract-new"
    if a.get("contract", type=int):
        return "contract-%s-%d" % (a.get("mode", "view"), a.get("contract", type=int))
    if a.get("new_cert"):
        return "cert-new"
    if a.get("cert", type=int):
        return "cert-%d" % a.get("cert", type=int)
    return None


@route("/subcontractors/<int:sub_id>")
@auth.login_required
def subcontractor_detail(sub_id):
    return render_sub_page(sub_id)


@route("/subcontractors/<int:sub_id>/edit", methods=["POST"])
@auth.login_required
def subcontractor_edit(sub_id):
    if not db.query("SELECT [SubcontractorID] FROM [Subcontractors] WHERE [SubcontractorID] = ?", (sub_id,)):
        flash("لم يتم العثور على المقاول.")
        return redirect(url_for("subcontractors_list"))
    name, cr, tax, bank, phone, notes = _form()
    if not name:
        flash("اسم المقاول مطلوب.")
    elif any(r["SubcontractorID"] != sub_id and norm_key(r["Name"]) == norm_key(name) for r in db.query("SELECT [SubcontractorID], [Name] FROM [Subcontractors]")):
        flash("يوجد مقاول آخر بهذا الاسم.")
    else:
        db.execute("UPDATE [Subcontractors] SET [Name] = ?, [CR] = ?, [TaxNumber] = ?, [BankAccount] = ?, [Phone] = ?, [Notes] = ? "
                   "WHERE [SubcontractorID] = ?", (name, cr, tax, bank, phone, notes, sub_id))
        import contract_gen
        contract_gen.save_party_extras("مقاول", sub_id, request.form)
        flash("تم حفظ بيانات المقاول.")
    return redirect(url_for("subcontractor_detail", sub_id=sub_id))


@route("/subcontractors/<int:sub_id>/report")
@auth.login_required
def subcontractor_report(sub_id):
    subs, contracts, certs, _p = cs.load()
    sub = subs.get(sub_id)
    if not sub:
        flash("لم يتم العثور على المقاول.")
        return redirect(url_for("subcontractors_list"))
    rows = []
    for c in sorted((c for c in contracts.values() if c.get("SubcontractorID") == sub_id), key=lambda c: c.get("CreatedAt") or datetime.min, reverse=True):
        f = cs.contract_figures(c, certs)
        for g in f["certs"]:
            g["Vat"] = g["VAT"]
        rows.append(dict(c, certificates=sorted(f["certs"], key=lambda g: (g.get("SiteName") or "", g["CertificateNo"] or 0)),
                         certified_total=f["due"], paid_total=f["paid"]))
    return render_template("subcontractor_report.html", sub=sub, contracts=rows, report_date=datetime.now().strftime("%Y-%m-%d"))
