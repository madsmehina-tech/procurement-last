# -*- coding: utf-8 -*-
"""مقارنة عروض الأسعار - part of the purchasing cycle, BEFORE the purchase order / the contract.
Steps: 1 the comparison · 2 upload the offers · 3 check what was read · 4 match the items · 5 compare and award."""
import io
import logging
import os
import re
import tempfile
from datetime import datetime

from flask import flash, redirect, render_template, request, send_file, url_for

import archive
import auth
import compare_service as cps
import cycle_files
import db
import write_lock
import report_xlsx
from order_services import official_supplier, supplier_project_choices
from routes import route
from routes.subcontractors import project_site_lists
from textnorm import clean_spaces, norm_key

log = logging.getLogger("po.error")
ALLOWED = (".pdf", ".xlsx", ".xls", ".docx", ".doc", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".md", ".html", ".htm")


def _user():
    return (auth.current_user() or {}).get("Username", "")


def _date(raw):
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime((raw or "").strip(), fmt)
        except ValueError:
            continue
    return None


def _num(raw):
    try:
        return float(str(raw).replace(",", "").strip()) if str(raw).strip() != "" else None
    except ValueError:
        return None


def _get(compare_id):
    c = cps.get(compare_id)
    if not c:
        flash("المقارنة غير موجودة.")
    return c


def parse_pasted(text):
    """Rows pasted from Excel (TAB) or typed with ; or |: name, unit, qty, price [, total]. Returns [{name, unit, qty, price}]."""
    out = []
    for raw in (text or "").splitlines():
        if not raw.strip():
            continue
        cells = [c.strip() for c in re.split(r"\t|;|\|", raw)]
        cells = [c for c in cells if c != ""] if len(cells) > 5 else cells
        if len(cells) < 3:
            continue
        nums = [_num(c) for c in cells]
        if len(cells) >= 4 and nums[2] is not None:                              # name, unit, qty, price [, total]
            name, unit, qty, price = cells[0], cells[1], nums[2], nums[3]
        elif len(cells) == 3 and nums[1] is not None and nums[2] is not None:    # name, qty, price
            name, unit, qty, price = cells[0], "", nums[1], nums[2]
        else:
            continue
        if qty is None and price is None:
            continue
        out.append({"name": name, "unit": unit, "qty": qty or 0, "price": price or 0})
    return out


OWN_NAMES = ("الهضب", "al-hadab", "alhadab", "al hadab", "hadab")
JUNK_WORDS = ("customer", "cutomer", "client", "attention", "attn", "quotation", "estimate", "invoice", "tel", "fax", "date", "ref", "page", "no.:")


def clean_read_party(read, kind="supplier"):
    """What the reader found as the supplier's name is only a SUGGESTION: an e-mail, a web address, our own company, a «Customer Name» line or a
    string of digits is NOT a supplier. Otherwise it is matched with the registered suppliers (a damaged name takes the registered spelling)."""
    read = clean_spaces(read)
    low = read.lower()
    if not read or len(read) < 3 or len(read) > 90:
        return ""
    if "@" in low or "http" in low or "www." in low or any(w in low for w in OWN_NAMES) or any(w in low for w in JUNK_WORDS):
        return ""
    if sum(ch.isdigit() for ch in read) > len(read) * 0.3:
        return ""
    if kind == "supplier":
        from order_services import suggest_supplier
        match, _score = suggest_supplier(read)
        return match or read
    return read


MAX_OCR_PAGES = 10


def _pages(raw):
    """«1-3» / «2» -> (first, last) pages to OCR (default: the first three - the price table is there, the rest are drawings and terms)."""
    m = re.fullmatch(r"\s*(\d+)\s*(?:-\s*(\d+))?\s*", raw or "")
    if not m:
        return (1, 3)
    a = max(int(m.group(1)), 1)
    return (a, min(max(a, int(m.group(2) or a)), a + MAX_OCR_PAGES - 1))          # at most MAX_OCR_PAGES pages: «1-9999» must not keep the server busy for hours


def _party_ids(kind, name):
    key = norm_key(name)
    if kind == "contractor":
        for s in db.query("SELECT [SubcontractorID], [Name] FROM [Subcontractors]"):
            if norm_key(s["Name"]) == key:
                return None, s["SubcontractorID"], s["Name"]
    else:
        for s in db.query("SELECT [SupplierID], [Name] FROM [Suppliers]"):
            if norm_key(s["Name"]) == key:
                return s["SupplierID"], None, s["Name"]
    return None, None, clean_spaces(name)


# --------------------------------------------------------------------------- the list
@route("/compare")
@auth.login_required
def compare_list():
    rows = db.query("SELECT * FROM [Comparisons]")
    quotes = db.query("SELECT * FROM [Quotes]")
    by = {}
    for q in quotes:
        by.setdefault(q["CompareID"], []).append(q)
    q_of = {q["QuoteID"]: q for q in quotes}
    view = [dict(c, offers=len(by.get(c["CompareID"], [])), winner=(q_of.get(c.get("WinnerQuoteID")) or {}).get("PartyName", "")) for c in sorted(rows, key=lambda c: c["CreatedAt"] or datetime.min, reverse=True)]
    return render_template("compare_list.html", rows=view, stats=cps.stats())


@route("/compare/new", methods=["GET", "POST"])
@auth.login_required
def compare_new():
    projects, sites_by_project = project_site_lists()
    if request.method == "POST":
        f = request.form
        try:
            cid = cps.create(f.get("title"), f.get("kind"), f.get("project"), f.get("site"), f.get("description"), _user())
        except cps.CompareError as e:
            flash(str(e))
            return render_template("compare_new.html", projects=projects, sites_by_project=sites_by_project, form=f, kinds=cps.KINDS)
        flash("أُنشئت المقارنة. الخطوة التالية: ارفع عروض الأسعار.")
        return redirect(url_for("compare_view", compare_id=cid))
    return render_template("compare_new.html", projects=projects, sites_by_project=sites_by_project, form={}, kinds=cps.KINDS)


# --------------------------------------------------------------------------- the workspace (steps 1-2)
@route("/compare/<int:compare_id>")
@auth.login_required
def compare_view(compare_id):
    c = _get(compare_id)
    if not c:
        return redirect(url_for("compare_list"))
    quotes = sorted(db.query("SELECT * FROM [Quotes] WHERE [CompareID] = ?", (compare_id,)), key=lambda q: q["QuoteID"])
    n_lines = {q["QuoteID"]: len(db.query("SELECT [LineID] FROM [QuoteLines] WHERE [QuoteID] = ?", (q["QuoteID"],))) for q in quotes}
    files = {q["QuoteID"]: archive.list_documents("QUOTE", q["QuoteID"]) for q in quotes}
    suppliers, _p = supplier_project_choices()
    contractors = [s["Name"] for s in db.query("SELECT [Name] FROM [Subcontractors]")]
    return render_template("compare_view.html", c=c, quotes=quotes, is_unnamed=cps.is_placeholder, n_lines=n_lines, files=files, suppliers=[s["SupplierName"] for s in suppliers], contractors=sorted(contractors),
                           ocr_ready=_ocr_ready(), ai_ready=_ai_ready(), step=_step(c, quotes), editable=c["Status"] not in (cps.ST_AWARDED, cps.ST_CANCELLED))


def _ocr_ready():
    try:
        import quote2po
        return bool(quote2po.ocr_available())
    except Exception:                                    # noqa
        return False


def _ai_ready():
    try:
        import quote2po
        return bool(quote2po.has_api_key())
    except Exception:                                    # noqa
        return False


def _step(c, quotes):
    if c["Status"] == cps.ST_AWARDED:
        return 5
    if not quotes:
        return 2
    return 3 if not c.get("Confirmed") else 5


@route("/compare/<int:compare_id>/quotes/new", methods=["POST"])
@auth.login_required
def compare_quote_add(compare_id):
    """A new offer: from a file (read by quote2po), from rows pasted from Excel, or an empty one to type."""
    c = _get(compare_id)
    if not c:
        return redirect(url_for("compare_list"))
    f = request.form
    kind = "contractor" if f.get("party_kind") == "contractor" else "supplier"
    upload = request.files.get("quote_file")
    pasted = f.get("paste", "")
    lines, res, engine, warnings, tmp = [], {}, "", [], None
    if upload and upload.filename:
        ext = os.path.splitext(upload.filename)[1].lower()
        if ext not in ALLOWED:
            flash("صيغة الملف غير مدعومة (المسموح: PDF، Word، Excel، صور).")
            return redirect(url_for("compare_view", compare_id=compare_id))
        fd, tmp = tempfile.mkstemp(suffix=ext)
        os.close(fd)
        upload.save(tmp)
        md = ""
        try:
            import quote2po
            import quote_read
            known = [x["SupplierName"] for x in supplier_project_choices()[0]]
            own = [n for n in OWN_NAMES if " " not in n or n == "al hadab"] + ["شركة الهضب للتجارة والمقاولات"]
            try:
                with write_lock.released():                   # the (possibly slow) reading must not freeze everybody else's saves
                    got = quote_read.read_offer(tmp, known=known, exclude=own, pages=_pages(f.get("ocr_pages")))
                md, engine, res = got["md"], got["engine"], got["res"]
                lines = got["items"]
                if engine == "ocr":
                    warnings.append("قُرئ الملف بالتعرف الضوئي (OCR): راجع كل اسم ورقم مع الملف الأصلي")
                warnings.extend(got["notes"])
            except Exception as e:                       # noqa - the AI step below may still read it
                log.exception("quote read failed")
                warnings.append("تعذّرت القراءة المباشرة للملف (%s)" % " ".join(str(e).split())[:110])
            if not lines and quote2po.has_api_key():
                try:                                     # the optional second reader: only when a key is configured and the direct reading found nothing
                    with write_lock.released():           # up to 2 minutes on the network
                        res = quote2po.ai_extract(tmp, md)
                    engine = "ai"
                    lines = res.get("items", []) if res else []
                    warnings.append("قُرئ العرض بالذكاء الاصطناعي لأن القراءة المباشرة لم تجد بنودًا: راجع كل سطر")
                except Exception as e:                   # noqa
                    log.exception("quote AI read failed")
                    warnings.append("تعذّرت قراءة الذكاء الاصطناعي (%s)" % " ".join(str(e).split())[:110])
            if not lines:
                warnings.append("أدخل البنود يدويًا أو الصق جدولًا من Excel")
        finally:
            try:
                os.remove(tmp)
            except OSError:
                pass
    elif pasted.strip():
        lines = parse_pasted(pasted)
        engine = "paste"
        if not lines:
            flash("لم أفهم الجدول الملصوق. الترتيب: البند، الوحدة، الكمية، سعر الوحدة (مفصولة بـ Tab أو ;).")
            return redirect(url_for("compare_view", compare_id=compare_id))
    typed = clean_spaces(f.get("party"))
    read_name = clean_read_party(res.get("supplier") if res else "", kind)
    party = typed or read_name
    if res and res.get("supplier") and not typed and not read_name:
        warnings.append("لم يُعرف اسم المورد من الملف (قُرئ «%s» وهو ليس اسم مورد)؛ العرض محفوظ باسم مؤقت: حدّد الاسم من صفحة المراجعة." % clean_spaces(res["supplier"])[:60])
    elif not party and (upload and upload.filename):
        warnings.append("لم يُعرف اسم المورد من الملف؛ العرض محفوظ باسم مؤقت: حدّد الاسم من صفحة المراجعة.")
    sid, subid, party = _party_ids(kind, party) if party else (None, None, "")
    if party and kind == "supplier" and not sid:
        party = official_supplier(party)
    try:
        qid = cps.add_quote(compare_id, kind, party, sid, subid, f.get("quote_no") or res.get("quote_no") or "", _date(f.get("quote_date")) or _date(res.get("quote_date")), f.get("vat", ""),
                            f.get("payment_terms", ""), int(f["delivery_days"]) if (f.get("delivery_days") or "").strip().isdigit() else None, _date(f.get("valid_until")), _num(f.get("declared_total")),
                            engine, upload.filename if (upload and upload.filename) else "", lines, _user())
    except cps.CompareError as e:
        flash(str(e))
        return redirect(url_for("compare_view", compare_id=compare_id))
    if upload and upload.filename:
        try:
            folder = archive.ensure_compare_folder(compare_id)
            upload.stream.seek(0)
            archive.store_document("QUOTE", qid, archive.compare_doc_folder_rel(folder, "عرض سعر أصلي"), "عرض سعر أصلي", upload, upload.filename, "", _user())
        except Exception as e:                           # noqa
            log.exception("quote archive failed")
            warnings.append("حُفظ العرض لكن تعذّر حفظ الملف الأصلي في الأرشيف (%s)." % " ".join(str(e).split())[:100])
    shown = party or cps.get_quote_name(qid)
    head = ("قُرئ %d بند من عرض «%s»." % (len(lines), shown)) if lines else "لم تُقرأ بنود من عرض «%s»." % shown
    flash(head + ((" " + " · ".join(warnings) + ".") if warnings else "") + " راجع الاسم والبنود قبل المطابقة.")
    return redirect(url_for("compare_quote", compare_id=compare_id, quote_id=qid))


# --------------------------------------------------------------------------- step 3: check what was read
@route("/compare/<int:compare_id>/quotes/<int:quote_id>")
@auth.login_required
def compare_quote(compare_id, quote_id):
    c = _get(compare_id)
    q = db.query("SELECT * FROM [Quotes] WHERE [QuoteID] = ? AND [CompareID] = ?", (quote_id, compare_id))
    if not c or not q:
        return redirect(url_for("compare_list"))
    lines = sorted(db.query("SELECT * FROM [QuoteLines] WHERE [QuoteID] = ?", (quote_id,)), key=lambda l: l["LineNo"] or 0)
    docs = archive.list_documents("QUOTE", quote_id)
    suppliers, _p = supplier_project_choices()
    return render_template("compare_quote.html", c=c, q=q[0], lines=lines, docs=docs, ocr_ready=_ocr_ready(), editable=c["Status"] not in (cps.ST_AWARDED, cps.ST_CANCELLED),
                           suppliers=[x["SupplierName"] for x in suppliers], contractors=sorted(x["Name"] for x in db.query("SELECT [Name] FROM [Subcontractors]")), unnamed=cps.is_placeholder(q[0]["PartyName"]))


@route("/compare/<int:compare_id>/quotes/<int:quote_id>/save", methods=["POST"])
@auth.login_required
def compare_quote_save(compare_id, quote_id):
    f = request.form
    lines = []
    for name, unit, qty, price in zip(f.getlist("name"), f.getlist("unit"), f.getlist("qty"), f.getlist("price")):
        if clean_spaces(name):
            lines.append({"name": name, "unit": unit, "qty": _num(qty) or 0, "price": _num(price) or 0})
    pasted = parse_pasted(f.get("paste", ""))
    try:
        kind = "contractor" if f.get("party_kind") == "contractor" else "supplier"
        if f.get("party") is not None:
            sid, subid, pname = _party_ids(kind, f.get("party"))
            if kind == "supplier" and not sid:
                pname = official_supplier(pname)
            cps.set_party(quote_id, kind, pname, sid, subid)
        cps.update_quote(quote_id, QuoteNo=(f.get("quote_no") or "")[:60], QuoteDate=_date(f.get("quote_date")), VatIncluded=f.get("vat") if f.get("vat") in ("yes", "no") else "",
                         PaymentTerms=(f.get("payment_terms") or "")[:150], DeliveryDays=int(f["delivery_days"]) if (f.get("delivery_days") or "").strip().isdigit() else None,
                         ValidUntil=_date(f.get("valid_until")), DeclaredTotal=_num(f.get("declared_total")), Notes=(f.get("notes") or "")[:255])
        cps.replace_lines(quote_id, lines + pasted)
        stats = cps.rematch(compare_id)
    except cps.CompareError as e:
        flash(str(e))
        return redirect(url_for("compare_quote", compare_id=compare_id, quote_id=quote_id))
    flash("حُفظ العرض (%d بند)%s." % (len(lines) + len(pasted), (" · أُضيف %d من اللصق" % len(pasted)) if pasted else ""))
    nxt = f.get("next")
    return redirect(url_for("compare_match", compare_id=compare_id) if nxt == "match" else url_for("compare_view", compare_id=compare_id))


@route("/compare/<int:compare_id>/quotes/<int:quote_id>/delete", methods=["POST"])
@auth.login_required
def compare_quote_delete(compare_id, quote_id):
    try:
        cps.delete_quote(quote_id)
        cps.rematch(compare_id) if db.query("SELECT [QuoteID] FROM [Quotes] WHERE [CompareID] = ?", (compare_id,)) else None
        flash("حُذف العرض من المقارنة (ملفه الأصلي باقٍ في الأرشيف).")
    except cps.CompareError as e:
        flash(str(e))
    return redirect(url_for("compare_view", compare_id=compare_id))


# --------------------------------------------------------------------------- step 4: match the items
@route("/compare/<int:compare_id>/match")
@auth.login_required
def compare_match(compare_id):
    c = _get(compare_id)
    if not c:
        return redirect(url_for("compare_list"))
    a = cps.analyze(compare_id)
    if not a["quotes"]:
        flash("ارفع عرضًا واحدًا على الأقل أولًا.")
        return redirect(url_for("compare_view", compare_id=compare_id))
    groups = []
    for r in a["rows"]:
        members = []
        for q in a["quotes"]:
            cell = r["cells"][q["QuoteID"]]
            if cell.get("present"):
                l = cell["line"]
                members.append({"line": l, "party": q["PartyName"], "review": cell["review"], "conf": l.get("Confidence") or 0, "locked": bool(l.get("Locked"))})
        groups.append({"item": r["item"], "name": r["name"], "unit": r["unit"], "required": r["required"], "fixed": r["qty_fixed"], "members": members, "missing": [q["PartyName"] for q in a["quotes"] if not r["cells"][q["QuoteID"]].get("present")]})
    return render_template("compare_match.html", c=c, a=a, groups=groups, editable=c["Status"] not in (cps.ST_AWARDED, cps.ST_CANCELLED), step=3 if not c.get("Confirmed") else 5)


@route("/compare/<int:compare_id>/match/save", methods=["POST"])
@auth.login_required
def compare_match_save(compare_id):
    f = request.form
    try:
        assignments = {k[len("line_"):]: v for k, v in f.items() if k.startswith("line_") and v}
        current = {str(l["LineID"]): str(l.get("ItemID")) for q in db.query("SELECT [QuoteID] FROM [Quotes] WHERE [CompareID] = ?", (compare_id,)) for l in db.query("SELECT * FROM [QuoteLines] WHERE [QuoteID] = ?", (q["QuoteID"],))}
        changed = {k: v for k, v in assignments.items() if current.get(k) != v}
        if changed:
            cps.assign(compare_id, changed, _user())
        cps.update_items(compare_id, names={k[len("name_"):]: v for k, v in f.items() if k.startswith("name_")}, units={k[len("unit_"):]: v for k, v in f.items() if k.startswith("unit_")},
                         qtys={k[len("qty_"):]: v for k, v in f.items() if k.startswith("qty_")})
        if f.get("action") == "rematch":
            s = cps.rematch(compare_id)
            flash("أُعيدت المطابقة للسطور غير المثبّتة: %d تلقائي · %d للمراجعة · %d بند جديد." % (s["auto"], s["review"], s["new"]))
            return redirect(url_for("compare_match", compare_id=compare_id))
        if f.get("action") == "confirm":
            cps.confirm_matching(compare_id, _user())
            flash("تم تأكيد المطابقة وحُفظت للمرات القادمة. هذا جدول المقارنة.")
            return redirect(url_for("compare_table", compare_id=compare_id))
        flash("حُفظت التعديلات.")
    except cps.CompareError as e:
        flash(str(e))
    return redirect(url_for("compare_match", compare_id=compare_id))


# --------------------------------------------------------------------------- step 5: the table, the decision, the award
@route("/compare/<int:compare_id>/table")
@auth.login_required
def compare_table(compare_id):
    c = _get(compare_id)
    if not c:
        return redirect(url_for("compare_list"))
    a = cps.analyze(compare_id)
    if not a["quotes"]:
        flash("ارفع عرضًا واحدًا على الأقل أولًا.")
        return redirect(url_for("compare_view", compare_id=compare_id))
    winner = next((s for s in a["summary"] if s["quote"]["QuoteID"] == c.get("WinnerQuoteID")), None)
    return render_template("compare_table.html", c=c, a=a, winner=winner, th=_thresholds(), docs=archive.list_documents("COMPARE", compare_id), editable=c["Status"] not in (cps.ST_AWARDED, cps.ST_CANCELLED))


def _thresholds():
    return {"green": cps.GREEN_BELOW, "red": cps.RED_ABOVE, "very_low": cps.VERY_LOW, "complete": cps.MIN_COMPLETE}


@route("/compare/<int:compare_id>/award", methods=["POST"])
@auth.login_required
def compare_award(compare_id):
    try:
        r = cps.award(compare_id, request.form.get("quote_id", type=int), request.form.get("reason", ""), _user())
    except cps.CompareError as e:
        flash(str(e))
        return redirect(url_for("compare_table", compare_id=compare_id))
    note = []
    try:                                                       # the report and the winning offer go to the archive
        folder = archive.ensure_compare_folder(compare_id)
        pdf = cycle_files._pdf_bytes(_print_html(compare_id))
        archive.store_generated("COMPARE", compare_id, archive.compare_doc_folder_rel(folder, "تقرير المقارنة"), "تقرير المقارنة", pdf, "تقرير مقارنة %s.pdf" % cps.get(compare_id)["Title"][:60], _user())
        note.append("حُفظ تقرير المقارنة في الأرشيف")
    except Exception as e:                                     # noqa - the award stands; the report can be saved by hand
        note.append("لم يُحفظ تقرير PDF تلقائيًا (%s) - احفظه بزر «حفظ التقرير في الأرشيف»" % " ".join(str(e).split())[:80])
    try:
        win_q = r["winner"]["quote"]["QuoteID"]
        for d in archive.list_documents("QUOTE", win_q):
            folder = archive.ensure_compare_folder(compare_id)
            archive.store_document("COMPARE", compare_id, archive.compare_doc_folder_rel(folder, "العرض الفائز"), "العرض الفائز", archive.abs_path(d["RelPath"]), d["OriginalName"], "", _user())
        note.append("ونُسخ العرض الفائز")
    except Exception:                                          # noqa
        pass
    flash("تمت الترسية على «%s» · الوفر مقابل متوسط باقي العروض %s ريال. %s." % (r["winner"]["name"], "{:,.2f}".format(r["savings_vs_avg"]), " · ".join(note)))
    return redirect(url_for("compare_table", compare_id=compare_id))


@route("/compare/<int:compare_id>/reopen", methods=["POST"])
@auth.login_required
def compare_reopen(compare_id):
    try:
        cps.reopen(compare_id)
        flash("أُلغيت الترسية وعادت المقارنة للمراجعة.")
    except cps.CompareError as e:
        flash(str(e))
    return redirect(url_for("compare_table", compare_id=compare_id))


@route("/compare/<int:compare_id>/next-ref", methods=["POST"])
@auth.login_required
def compare_next_ref(compare_id):
    ref = clean_spaces(request.form.get("ref"))
    c = cps.get(compare_id)
    if c and ref:
        cps.link_next(compare_id, c.get("NextKind") or "أمر شراء", ref)
        flash("رُبطت الترسية بـ %s %s." % (c.get("NextKind") or "أمر شراء", ref))
    return redirect(url_for("compare_table", compare_id=compare_id))


@route("/compare/<int:compare_id>/delete", methods=["POST"])
@auth.login_required
def compare_delete(compare_id):
    try:
        cps.delete_comparison(compare_id)
        flash("حُذفت المقارنة (ملفاتها الأصلية باقية في الأرشيف).")
    except cps.CompareError as e:
        flash(str(e))
        return redirect(url_for("compare_table", compare_id=compare_id))
    return redirect(url_for("compare_list"))


# --------------------------------------------------------------------------- reports
def _print_html(compare_id):
    c = cps.get(compare_id)
    a = cps.analyze(compare_id)
    winner = next((s for s in a["summary"] if s["quote"]["QuoteID"] == c.get("WinnerQuoteID")), None)
    return render_template("compare_print.html", c=c, a=a, winner=winner, th=_thresholds(), now=datetime.now().strftime("%Y-%m-%d %H:%M"))


@route("/compare/<int:compare_id>/print")
@auth.login_required
def compare_print(compare_id):
    if not cps.get(compare_id):
        return redirect(url_for("compare_list"))
    return _print_html(compare_id)


@route("/compare/<int:compare_id>/report.pdf")
@auth.login_required
def compare_pdf(compare_id):
    if not cps.get(compare_id):
        return redirect(url_for("compare_list"))
    try:
        return send_file(io.BytesIO(cycle_files._pdf_bytes(_print_html(compare_id))), as_attachment=True, download_name="comparison_%s.pdf" % compare_id, mimetype="application/pdf")
    except cycle_files.pdfgen.PdfError as e:
        flash("تعذّر إنشاء PDF تلقائيًا (%s) - استخدم «طباعة» ثم «حفظ كـ PDF» من المتصفح." % e)
        return redirect(url_for("compare_table", compare_id=compare_id))


@route("/compare/<int:compare_id>/report/archive", methods=["POST"])
@auth.login_required
def compare_report_archive(compare_id):
    c = cps.get(compare_id)
    if not c:
        return redirect(url_for("compare_list"))
    try:
        folder = archive.ensure_compare_folder(compare_id)
        archive.store_generated("COMPARE", compare_id, archive.compare_doc_folder_rel(folder, "تقرير المقارنة"), "تقرير المقارنة", cycle_files._pdf_bytes(_print_html(compare_id)), "تقرير مقارنة %s.pdf" % c["Title"][:60], _user())
        flash("حُفظ تقرير المقارنة (PDF) في الأرشيف.")
    except Exception as e:                                     # noqa
        flash("لم يُحفظ التقرير: %s" % " ".join(str(e).split())[:160])
    return redirect(url_for("compare_table", compare_id=compare_id))


@route("/compare/<int:compare_id>/export.xlsx")
@auth.login_required
def compare_xlsx(compare_id):
    c = cps.get(compare_id)
    if not c:
        return redirect(url_for("compare_list"))
    a = cps.analyze(compare_id)
    qs = a["quotes"]
    head = ["البند", "الوحدة", "الكمية المطلوبة"] + sum(([q["PartyName"] + " - سعر الوحدة", q["PartyName"] + " - الإجمالي"] for q in qs), []) + ["متوسط العروض", "آخر سعر شراء", "متوسط تاريخي"]
    rows = []
    for r in a["rows"]:
        row = [r["name"], r["unit"], r["required"]]
        for q in qs:
            cell = r["cells"][q["QuoteID"]]
            row += [cell["price"], cell["total"]] if cell.get("present") else ["غير موجود", "غير موجود"]
        po = r["history"]["po"] or {}
        rows.append(row + [r["avg_offer"], po.get("last") or "", r["hist_avg"] or ""])
    sums = [[s["name"], s["quote"].get("QuoteNo") or "", {"yes": "شاملة", "no": "غير شاملة"}.get(s["quote"].get("VatIncluded"), "غير محدد"), s["quoted"], s["normalized"], s["completing"], s["adjusted"],
             s["complete"], len(s["missing"]), s["red"], " | ".join(s["flags"])] for s in a["summary"]]
    maps = [[q["PartyName"], l["RawName"], l.get("Unit") or "", l.get("Qty") or 0, l.get("UnitPrice") or 0, next((i["Name"] for i in a["rows"] and [r["item"] for r in a["rows"]] if i["ItemID"] == l.get("ItemID")), ""), l.get("Confidence") or 0]
            for q in qs for l in sorted(db.query("SELECT * FROM [QuoteLines] WHERE [QuoteID] = ?", (q["QuoteID"],)), key=lambda l: l["LineNo"] or 0)]
    hist = [[r["name"], r["unit"], (r["history"]["po"] or {}).get("last") or "", (r["history"]["po"] or {}).get("avg") or "", (r["history"]["po"] or {}).get("max") or "", (r["history"]["po"] or {}).get("min") or "",
             (r["history"]["po"] or {}).get("count") or 0, (r["history"]["prior"] or {}).get("avg") or ""] for r in a["rows"]]
    buf = report_xlsx.build([
        {"title": "المقارنة", "band": "مقارنة عروض الأسعار: " + c["Title"], "headers": head, "rows": rows, "widths": (36, 10, 12) + (16, 16) * len(qs) + (14, 14, 14), "money": tuple(range(4, len(head) + 1)), "table": "tblCompare"},
        {"title": "العروض", "band": "ملخص العروض", "headers": ["العرض", "الرقم", "الضريبة", "الإجمالي كما قُدِّم", "الإجمالي (الكمية المطلوبة، قبل الضريبة)", "تكلفة استكمال الناقص", "الإجمالي المعدّل", "الاكتمال %", "بنود ناقصة", "أسعار حمراء", "تنبيهات"],
         "rows": sums, "widths": (34, 14, 14, 18, 22, 18, 18, 12, 12, 12, 60), "money": (4, 5, 6, 7), "table": "tblOffers"},
        {"title": "مطابقة البنود", "band": "كل سطر كما كُتب والبند الموحّد", "headers": ["العرض", "البند كما كُتب", "الوحدة", "الكمية", "السعر", "البند الموحّد", "ثقة المطابقة"], "rows": maps, "widths": (30, 44, 10, 12, 14, 36, 12), "money": (5,), "table": "tblMatches"},
        {"title": "الأسعار التاريخية", "band": "سجل الشراء لكل بند", "headers": ["البند", "الوحدة", "آخر سعر", "المتوسط", "الأعلى", "الأقل", "عدد مرات الشراء", "متوسط عروض سابقة"], "rows": hist, "widths": (36, 10, 14, 14, 14, 14, 14, 16), "money": (3, 4, 5, 6, 8), "table": "tblHistory"}],
        "%s · %s" % (c.get("ProjectName") or "", c.get("SiteName") or ""))
    return send_file(buf, as_attachment=True, download_name="comparison_%s.xlsx" % compare_id, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@route("/compare/template.xlsx")
@auth.login_required
def compare_template():
    """A ready Excel template for typing / importing an offer when the file cannot be read automatically."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "عرض سعر"
    ws.sheet_view.rightToLeft = True
    ws.append(["اسم المورد: ..."])
    ws.append(["رقم العرض: ..."])
    ws.append([])
    ws.append(["م", "البيان", "الوحدة", "الكمية", "سعر الوحدة", "الإجمالي"])
    for c in ws[4]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", start_color="14323A")
    ws.append([1, "مثال: كابل نحاس 4×16", "متر", 100, 38, "=D5*E5"])
    for col, w in zip("ABCDEF", (6, 46, 12, 12, 14, 16)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="quote_template.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# --------------------------------------------------------------------------- the next documents
@route("/compare/<int:compare_id>/to-order")
@auth.login_required
def compare_to_order(compare_id):
    """The purchase order form, filled with the winning offer (the original quote is attached to the order when it is saved)."""
    from routes import orders as orders_routes
    try:
        pre = cps.po_prefill(compare_id)
    except cps.CompareError as e:
        flash(str(e))
        return redirect(url_for("compare_table", compare_id=compare_id))
    form = dict(pre["form"])
    docs = archive.list_documents("QUOTE", pre["quote"]["QuoteID"])
    if docs:
        try:
            token = archive.save_pending(_FileLike(archive.abs_path(docs[0]["RelPath"]), docs[0]["OriginalName"]))
            form["quote_file_path"], form["quote_file_name"] = token, docs[0]["OriginalName"]
        except Exception:                                      # noqa - the order can still be created; the quote is attached by hand
            pass
    warnings = ["المقارنة: بنود غير موجودة في العرض الفائز ولم تُضف: " + "، ".join(pre["missing"])] if pre["missing"] else []
    warnings.append("الأسعار هنا قبل الضريبة وبالكمية المطلوبة في المقارنة؛ راجع قبل الحفظ.")
    return render_template("order_new.html", **orders_routes._new_context(form, warnings=warnings))


class _FileLike:
    """A file on disk that archive.save_pending can read like an upload."""

    def __init__(self, path, name):
        self.filename, self._path = name, path
        self.stream = open(path, "rb")

    def save(self, dst):
        with open(dst, "wb") as out:
            self.stream.seek(0)
            out.write(self.stream.read())

    def read(self, *a):
        return self.stream.read(*a)


@route("/compare/<int:compare_id>/to-contract")
@auth.login_required
def compare_to_contract(compare_id):
    """The contractor's page with the contract window open and filled (project, value, notes)."""
    from routes import subcontractors as sub_routes
    c = cps.get(compare_id)
    q = db.query("SELECT * FROM [Quotes] WHERE [QuoteID] = ?", (c.get("WinnerQuoteID"),)) if c else []
    if not q or not q[0].get("SubcontractorID"):
        flash("العرض الفائز لمقاول غير مسجّل في النظام. سجّله أولًا من «مقاولو الباطن» ثم كرّر.")
        return redirect(url_for("compare_table", compare_id=compare_id))
    a = cps.analyze(compare_id)
    win = next(s for s in a["summary"] if s["quote"]["QuoteID"] == q[0]["QuoteID"])
    form = {"project": c.get("ProjectName") or "", "site_name": c.get("SiteName") or "", "contract_value": "%.2f" % (win["adjusted"] or win["normalized"]),
            "notes": "بموجب مقارنة عروض الأسعار: %s" % c["Title"]}
    flash("ملأنا نافذة العقد من العرض الفائز. بعد الحفظ اربط العقد بالترسية من صفحة المقارنة.")
    return sub_routes.render_sub_page(q[0]["SubcontractorID"], open_dialog="contract-new", form=form)
