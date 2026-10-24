# -*- coding: utf-8 -*-
"""Suppliers: directory with filters / sorting / Excel & PDF export, supplier file (data, representative, papers, orders).

The suppliers table is fed automatically by purchase orders (masterdata.resolve_supplier): a supplier typed on an order
appears here by itself; this page completes its data and papers."""
import io
from datetime import datetime

from flask import flash, redirect, render_template, request, send_file, url_for

import archive
import ui_helpers
import auth
import db
import pdfgen
from routes import route
from textnorm import clean_spaces, contains, norm_key

SLOTS_TOTAL = 4        # commercial register, tax certificate, catalog, other files
SORTS = {"name": "الاسم", "orders": "الأكثر أوامر", "total": "الأعلى مشتريات", "last": "أحدث تعامل", "files": "الأوراق الأقل اكتمالًا"}
DOCS = {"": "كل حالات الأوراق", "complete": "أوراق مكتملة", "incomplete": "أوراق ناقصة", "expiring": "أوراق منتهية أو قاربت"}
ACTIVITY = {"": "الكل", "active": "تعامل معه (له أوامر)", "inactive": "بلا أوامر بعد"}


def _order_stats():
    """{SupplierID or name key: [orders, total, last order date, projects]} - by ID when the order has one."""
    stats = {}
    for r in db.query("SELECT [SupplierID], [SupplierName], [TotalAmount], [OrderDate], [Project] FROM [PurchaseOrders]"):
        key = r.get("SupplierID") or norm_key(r["SupplierName"])
        s = stats.setdefault(key, [0, 0.0, None, set()])
        s[0] += 1
        s[1] += r["TotalAmount"] or 0
        if r.get("OrderDate") and (not s[2] or r["OrderDate"] > s[2]):
            s[2] = r["OrderDate"]
        if r.get("Project"):
            s[3].add(r["Project"])
    return stats


def _doc_state():
    out = {}
    for d in db.query("SELECT [EntityKey], [DocType], [IsCurrent], [ExpiryDate] FROM [Documents] WHERE [EntityType] = ?", ("SUPPLIER",)):
        if not d["IsCurrent"]:
            continue
        slot = d["DocType"] if d["DocType"] in archive.SUPPLIER_SLOTS else "ملفات أخرى"
        st = out.setdefault(str(d["EntityKey"]), {"slots": set(), "worst": None})
        st["slots"].add(slot)
        state = archive.expiry_state(d.get("ExpiryDate"))
        if state and state[0] in ("expired", "soon") and (not st["worst"] or state[0] == "expired"):
            st["worst"] = state
    return out


def _rows():
    a = request.args
    q, sort, f_project = a.get("q", "").strip(), a.get("sort", "name"), a.get("project", "").strip()
    f_docs, f_act = a.get("docs", ""), a.get("activity", "")
    stats, docs = _order_stats(), _doc_state()
    rows, projects = [], set()
    for r in db.query("SELECT * FROM [Suppliers]"):
        st = stats.get(r["SupplierID"]) or stats.get(norm_key(r["Name"])) or [0, 0.0, None, set()]
        if r["SupplierID"] in stats and norm_key(r["Name"]) in stats:                       # orders saved before the link existed
            old = stats[norm_key(r["Name"])]
            st = [st[0] + old[0], st[1] + old[1], max([d for d in (st[2], old[2]) if d] or [None]), st[3] | old[3]]
        projects |= st[3]
        d = docs.get(str(r["SupplierID"]), {"slots": set(), "worst": None})
        r = dict(r, orders=st[0], total=st[1], last=st[2], projects=st[3], files=len(d["slots"]), worst=d["worst"])
        if q and not any(contains(r.get(k) or "", q) for k in ("Name", "CR", "TaxNumber", "RepName", "Phone", "RepEmail")):
            continue
        if f_project and f_project not in r["projects"]:
            continue
        if f_docs == "complete" and r["files"] < SLOTS_TOTAL or f_docs == "incomplete" and r["files"] >= SLOTS_TOTAL or f_docs == "expiring" and not r["worst"]:
            continue
        if f_act == "active" and not r["orders"] or f_act == "inactive" and r["orders"]:
            continue
        rows.append(r)
    keys = {"name": (lambda r: norm_key(r["Name"]), False), "orders": (lambda r: r["orders"], True), "total": (lambda r: r["total"], True),
            "last": (lambda r: r["last"] or datetime.min, True), "files": (lambda r: r["files"], False)}
    k, rev = keys.get(sort, keys["name"])
    rows.sort(key=k, reverse=rev)
    kpi = {"count": len(rows), "total": sum(r["total"] for r in rows), "incomplete": sum(1 for r in rows if r["files"] < SLOTS_TOTAL),
           "expiring": sum(1 for r in rows if r["worst"])}
    filters = dict(q=q, sort=sort, project=f_project, docs=f_docs, activity=f_act)
    labels = [x for x in (q and "بحث: " + q, f_project and "المشروع: " + f_project, f_docs and DOCS[f_docs], f_act and ACTIVITY[f_act]) if x]
    return rows, kpi, filters, sorted(projects), " · ".join(labels) or "كل الموردين"


@route("/suppliers")
@auth.login_required
def suppliers_list():
    rows, kpi, f, projects, label = _rows()
    return render_template("suppliers_list.html", rows=rows, kpi=kpi, f=f, projects=projects, sorts=SORTS, docs_opts=DOCS,
                           act_opts=ACTIVITY, slots_total=SLOTS_TOTAL, filtered=any(v for k, v in f.items() if k != "sort"))


@route("/suppliers/export.xlsx")
@auth.login_required
def suppliers_export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    rows, kpi, _f, _p, label = _rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "الموردون"
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    navy, line = PatternFill("solid", start_color="1E293B"), Side(style="thin", color="E2E8F0")
    heads = ["#", "المورد", "اسم المندوب", "جوال المندوب", "البريد", "السجل التجاري", "الرقم الضريبي", "أوامر الشراء", "إجمالي المشتريات", "آخر أمر", "الأوراق"]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(heads))
    ws["A1"] = "سجل الموردين — شركة الهضب للتجارة والمقاولات"
    ws["A1"].font, ws["A1"].fill = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF"), navy
    ws["A1"].alignment = Alignment(horizontal="right", vertical="center", indent=1, readingOrder=2)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(heads))
    ws["A2"] = "%s  ·  %d مورد  ·  إجمالي المشتريات %s ريال  ·  صدر في %s" % (label, kpi["count"], "{:,.0f}".format(kpi["total"]), datetime.now().strftime("%Y-%m-%d %H:%M"))
    ws["A2"].font, ws["A2"].fill = Font(name="Segoe UI", size=10, color="CBD5E1"), navy
    ws["A2"].alignment = Alignment(horizontal="right", vertical="center", indent=1, readingOrder=2)
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 30, 20
    for j, h in enumerate(heads, 1):
        c = ws.cell(4, j, h)
        c.font, c.fill = Font(name="Segoe UI", bold=True, color="FFFFFF"), PatternFill("solid", start_color="29595E")
        c.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)
    for i, r in enumerate(rows, 1):
        vals = [i, r["Name"], r.get("RepName") or "", r.get("Phone") or "", r.get("RepEmail") or "", r.get("CR") or "", r.get("TaxNumber") or "",
                r["orders"], r["total"], r["last"].strftime("%Y-%m-%d") if r["last"] else "", "%d / %d" % (r["files"], SLOTS_TOTAL)]
        for j, v in enumerate(vals, 1):
            c = ws.cell(4 + i, j, v)
            c.font = Font(name="Segoe UI", size=10)
            c.border = Border(bottom=line)
            c.alignment = Alignment(horizontal="right" if j in (2, 3, 5) else "center", readingOrder=2)
            if j == 9:
                c.number_format = "#,##0.00"
            if i % 2 == 0:
                c.fill = PatternFill("solid", start_color="F8FAFC")
    last = 4 + len(rows)
    t = ws.cell(last + 1, 8, sum(r["orders"] for r in rows)), ws.cell(last + 1, 9, kpi["total"])
    ws.cell(last + 1, 2, "الإجمالي").font = Font(name="Segoe UI", bold=True)
    for c in t:
        c.font, c.number_format = Font(name="Segoe UI", bold=True), "#,##0"
    for j, w in enumerate((5, 34, 22, 15, 26, 14, 16, 12, 18, 12, 10), 1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:%s%d" % (chr(64 + len(heads)), max(5, last))
    ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", 9
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"
    ws.oddFooter.center.text = "سجل الموردين — صفحة &P من &N"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="suppliers_%s.xlsx" % datetime.now().strftime("%Y%m%d"),
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@route("/suppliers/export.pdf")
@auth.login_required
def suppliers_export_pdf():
    rows, kpi, _f, _p, label = _rows()
    html = render_template("suppliers_report.html", rows=rows, kpi=kpi, label=label, slots_total=SLOTS_TOTAL,
                           now=datetime.now().strftime("%Y-%m-%d %H:%M"), logo_src=_logo())
    if pdfgen.available():
        import os
        import tempfile
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        try:
            pdfgen.html_to_pdf(html, path)
            with open(path, "rb") as fh:
                data = fh.read()
            return send_file(io.BytesIO(data), as_attachment=True, download_name="suppliers_%s.pdf" % datetime.now().strftime("%Y%m%d"),
                             mimetype="application/pdf")
        except pdfgen.PdfError as e:
            flash("تعذّر إنشاء PDF تلقائيًا (%s) - استخدم «طباعة / حفظ PDF» من هذه الصفحة." % e)
        finally:
            try:
                os.remove(path)
            except OSError:
                pass
    return html                                              # no browser: the same report, printable from the browser


def _logo():
    import base64
    import os
    import config
    try:
        with open(os.path.join(config.BASE_DIR, "static", "company_logo.png"), "rb") as f:
            return "data:image/png;base64," + base64.b64encode(f.read()).decode("ascii")
    except OSError:
        return ""


def _form_values():
    f = request.form
    return (clean_spaces(f.get("name", "")), f.get("cr", "").strip(), f.get("tax_number", "").strip(), clean_spaces(f.get("rep_name", "")),
            f.get("phone", "").strip(), f.get("rep_email", "").strip(), f.get("bank_account", "").strip(), f.get("notes", "").strip())


@route("/suppliers/new", methods=["GET", "POST"])
@auth.login_required
def supplier_new():
    if request.method == "GET":
        return render_template("supplier_form.html", supplier=None)
    name, cr, tax, rep, phone, email, bank, notes = _form_values()
    if not name:
        flash("اسم المورد مطلوب.")
        return render_template("supplier_form.html", supplier=request.form)
    if any(norm_key(r["Name"]) == norm_key(name) for r in db.query("SELECT [Name] FROM [Suppliers]")):
        flash("يوجد مورد بهذا الاسم (أو اسم مطابق له). افتح ملفه بدل تكراره.")
        return render_template("supplier_form.html", supplier=request.form)
    supplier_id = db.execute_insert(
        "INSERT INTO [Suppliers] ([Name], [CR], [TaxNumber], [RepName], [Phone], [RepEmail], [BankAccount], [Notes], [CreatedAt]) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)", (name, cr, tax, rep, phone, email, bank, notes, datetime.now()))
    archive.ensure_supplier_folder(supplier_id)
    flash("تم إضافة المورد. ارفع أوراقه من قسم «الملفات» أدناه.")
    return redirect(url_for("supplier_detail", supplier_id=supplier_id))


@route("/suppliers/<int:supplier_id>")
@auth.login_required
def supplier_detail(supplier_id):
    rows = db.query("SELECT * FROM [Suppliers] WHERE [SupplierID] = ?", (supplier_id,))
    if not rows:
        flash("لم يتم العثور على المورد.")
        return redirect(url_for("suppliers_list"))
    s = rows[0]
    key = norm_key(s["Name"])
    orders = sorted((o for o in db.query("SELECT [PONumber], [OrderDate], [Project], [Location], [TotalAmount], [CreatedAt], [SupplierName], [SupplierID] "
                                         "FROM [PurchaseOrders]") if o.get("SupplierID") == supplier_id or norm_key(o["SupplierName"]) == key),
                    key=lambda o: o["CreatedAt"] or datetime.min, reverse=True)
    slots = archive.panel("SUPPLIER", supplier_id, archive.SUPPLIER_DOC_TYPES)
    by_type = {sl["type"]: sl for sl in slots}
    others = sorted((d for t in archive.SUPPLIER_OTHER_TYPES for d in by_type[t]["current"]),
                    key=lambda d: d["UploadedAt"] or datetime.min, reverse=True)
    filled = sum(1 for t in archive.SUPPLIER_SLOTS if by_type[t]["current"]) + (1 if others else 0)
    folder = s.get("ArchiveFolder") or archive.supplier_root_rel(s["Name"])
    next_actions = ui_helpers.paper_actions(ui_helpers.paper_alerts(slots, archive.SUPPLIER_SLOTS))
    if not orders:
        next_actions.append({"level": "info", "tab": "orders", "text": "لا توجد أوامر شراء لهذا المورد بعد"})
    return render_template(
        "supplier_detail.html", supplier=s, next_actions=next_actions, orders=orders[:8], orders_count=len(orders),
        total=sum(o["TotalAmount"] or 0 for o in orders), slots=by_type, others=others, filled=filled,
        slots_total=SLOTS_TOTAL, folder=folder + "/" + archive.SUPPLIER_PAPERS, network=archive.network_path(folder + "/" + archive.SUPPLIER_PAPERS),
        other_types=archive.SUPPLIER_OTHER_TYPES, slot_names=archive.SUPPLIER_SLOTS,
        projects=sorted({o["Project"] for o in orders if o.get("Project")}),
    )


@route("/suppliers/<int:supplier_id>/edit", methods=["POST"])
@auth.login_required
def supplier_edit(supplier_id):
    if not db.query("SELECT [SupplierID] FROM [Suppliers] WHERE [SupplierID] = ?", (supplier_id,)):
        flash("لم يتم العثور على المورد.")
        return redirect(url_for("suppliers_list"))
    name, cr, tax, rep, phone, email, bank, notes = _form_values()
    if not name:
        flash("اسم المورد مطلوب.")
    elif any(r["SupplierID"] != supplier_id and norm_key(r["Name"]) == norm_key(name) for r in db.query("SELECT [SupplierID], [Name] FROM [Suppliers]")):
        flash("يوجد مورد آخر بهذا الاسم.")
    else:
        db.execute("UPDATE [Suppliers] SET [Name] = ?, [CR] = ?, [TaxNumber] = ?, [RepName] = ?, [Phone] = ?, [RepEmail] = ?, [BankAccount] = ?, "
                   "[Notes] = ? WHERE [SupplierID] = ?", (name, cr, tax, rep, phone, email, bank, notes, supplier_id))
        db.execute("UPDATE [PurchaseOrders] SET [SupplierName] = ? WHERE [SupplierID] = ?", (name, supplier_id))   # its orders show the new name
        import contract_gen
        contract_gen.save_party_extras("مورد", supplier_id, request.form)
        flash("تم حفظ بيانات المورد.")
    return redirect(url_for("supplier_detail", supplier_id=supplier_id))
