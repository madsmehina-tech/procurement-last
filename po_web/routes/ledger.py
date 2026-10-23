# -*- coding: utf-8 -*-
"""The contractor's financial file: the movements typed by hand («الدفعات» section of the contractor's page) and the account
statement (كشف الحساب) with its print, Excel and PDF. See ledger_service for the rules."""
import io
import logging
from datetime import datetime

from flask import flash, redirect, render_template, request, send_file, url_for

import archive
import auth
import cycle_files
import ledger_service as ls
import payment_files
from routes import route


def _user():
    return (auth.current_user() or {}).get("Username", "")


def _day(raw):
    try:
        return datetime.strptime((raw or "").strip(), "%Y-%m-%d")
    except ValueError:
        return None


def _back(sub_id):
    return redirect(url_for("subcontractor_detail", sub_id=sub_id) + "#ledger")


# --------------------------------------------------------------------------- the movements typed by hand
@route("/subcontractors/<int:sub_id>/ledger/new", methods=["POST"])
@auth.login_required
def ledger_add(sub_id):
    f = request.form
    try:
        amount = float((f.get("amount") or "").replace(",", "") or 0)
    except ValueError:
        flash("المبلغ يجب أن يكون رقمًا.")
        return _back(sub_id)
    files = [x for x in request.files.getlist("files") if x and x.filename]
    bad = payment_files.bad_file(files)
    if bad:
        flash("صيغة الملف غير مدعومة: %s (المسموح: PDF، صور، Word، Excel) - لم تُحفظ الحركة." % bad)
        return _back(sub_id)
    try:
        entry_id = ls.add_entry(sub_id, f.get("kind"), f.get("category"), f.get("description"), amount, _day(f.get("entry_date")) or datetime.now(),
                                f.get("project"), f.get("site"), f.get("status", "executed") == "executed", f.get("bank_ref"), _user())
    except ls.LedgerError as e:
        flash(str(e))
        return _back(sub_id)
    except Exception as e:                                   # noqa
        logging.getLogger("po.error").exception("ledger save failed")
        flash("تعذّر حفظ الحركة ولم يتغير شيء. السبب: %s" % " ".join(str(e).split())[:200])
        return _back(sub_id)
    pending = f.get("category") == ls.CAT_PAYMENT and f.get("status") == "pending"
    saved = ""
    if files:
        try:
            saved = " - رُفع %d ملف إلى أرشيف الدفعات والحركات المالية" % payment_files.save_entry_files(entry_id, files, f.get("file_type"), f.get("description"), _user())
        except Exception as e:                                   # noqa - the movement is saved; only the file failed
            logging.getLogger("po.error").exception("ledger file failed")
            saved = " - لكن تعذّر رفع الملف (%s)؛ يمكن رفعه من الحركة" % " ".join(str(e).split())[:120]
    flash(("سُجّلت الدفعة ولم تُنفَّذ بعد؛ لن تظهر في كشف الحساب حتى «تم التحويل»، ويمكن إدراجها في دورة صرف" if pending else "سُجّلت الحركة وظهرت في كشف الحساب") + saved + ".")
    return _back(sub_id)


@route("/ledger/<int:entry_id>/files", methods=["POST"])
@auth.login_required
def ledger_files(entry_id):
    """Adds files (receipt, invoice ...) to a movement that already exists."""
    rows = ls.db.query("SELECT [SubcontractorID] FROM [LedgerEntries] WHERE [EntryID] = ?", (entry_id,))
    if not rows:
        flash("الحركة غير موجودة.")
        return redirect(url_for("subcontractors_list"))
    files = [x for x in request.files.getlist("files") if x and x.filename]
    if not files:
        flash("اختر ملفًا أولًا.")
    else:
        try:
            n = payment_files.save_entry_files(entry_id, files, request.form.get("file_type"), request.form.get("note", ""), _user())
            flash("رُفع %d ملف إلى أرشيف الحركة." % n)
        except Exception as e:                                   # noqa
            flash(" ".join(str(e).split())[:200])
    return _back(rows[0]["SubcontractorID"])


@route("/ledger/<int:entry_id>/execute", methods=["POST"])
@auth.login_required
def ledger_execute(entry_id):
    try:
        sub_id = ls.execute_entry(entry_id, _day(request.form.get("pay_date")) or datetime.now(), request.form.get("bank_ref"), _user())
    except ls.LedgerError as e:
        flash(str(e))
        return redirect(url_for("subcontractors_list"))
    archive.promote_entry_files(entry_id)                        # the files uploaded while it was only recorded move to the payment's folder
    flash("سُجّل «تم التحويل» وظهرت الدفعة في كشف الحساب.")
    return _back(sub_id)


@route("/ledger/<int:entry_id>/delete", methods=["POST"])
@auth.login_required
def ledger_delete(entry_id):
    try:
        sub_id = ls.delete_entry(entry_id)
    except ls.LedgerError as e:
        flash(str(e))
        return redirect(url_for("subcontractors_list"))
    flash("حُذفت الحركة.")
    return _back(sub_id)


# --------------------------------------------------------------------------- the statement
def _read(sub_id):
    a = request.args
    d_from, d_to, project, site = _day(a.get("date_from")), _day(a.get("date_to")), a.get("project", "").strip(), a.get("site", "").strip()
    data = ls.statement(sub_id, d_from, d_to, project, site)
    if not data["sub"]:
        return None, None
    period = []
    if a.get("date_from"):
        period.append("من %s" % a.get("date_from"))
    if a.get("date_to"):
        period.append("إلى %s" % a.get("date_to"))
    if project:
        period.append("مشروع: " + project)
    if site:
        period.append("موقع: " + site)
    return data, dict(date_from=a.get("date_from", ""), date_to=a.get("date_to", ""), project=project, site=site, label=" · ".join(period) or "كل الحركات")


@route("/subcontractors/<int:sub_id>/statement")
@auth.login_required
def statement_view(sub_id):
    data, f = _read(sub_id)
    if not data:
        flash("لم يتم العثور على المقاول.")
        return redirect(url_for("subcontractors_list"))
    return render_template("statement.html", d=data, f=f, sub=data["sub"])


@route("/subcontractors/<int:sub_id>/statement/print")
@auth.login_required
def statement_print(sub_id):
    data, f = _read(sub_id)
    if not data:
        return redirect(url_for("subcontractors_list"))
    return render_template("statement_print.html", d=data, f=f, sub=data["sub"], now=datetime.now().strftime("%Y-%m-%d %H:%M"))


@route("/subcontractors/<int:sub_id>/statement/export.pdf")
@auth.login_required
def statement_pdf(sub_id):
    data, f = _read(sub_id)
    if not data:
        return redirect(url_for("subcontractors_list"))
    html = render_template("statement_print.html", d=data, f=f, sub=data["sub"], now=datetime.now().strftime("%Y-%m-%d %H:%M"))
    try:
        pdf = cycle_files._pdf_bytes(html)
    except cycle_files.pdfgen.PdfError as e:
        flash("تعذّر إنشاء PDF تلقائيًا (%s) - استخدم «طباعة» ثم «حفظ كـ PDF» من المتصفح." % e)
        return redirect(url_for("statement_view", sub_id=sub_id, **{k: v for k, v in request.args.items()}))
    return send_file(io.BytesIO(pdf), as_attachment=True, download_name="statement_%s_%s.pdf" % (sub_id, datetime.now().strftime("%Y%m%d")), mimetype="application/pdf")


@route("/subcontractors/<int:sub_id>/statement/export.xlsx")
@auth.login_required
def statement_xlsx(sub_id):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    data, f = _read(sub_id)
    if not data:
        return redirect(url_for("subcontractors_list"))
    wb = Workbook()
    ws = wb.active
    ws.title = "كشف الحساب"
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    navy, line = PatternFill("solid", start_color="14323A"), Side(style="thin", color="E2E8F0")
    ws.merge_cells("A1:F1")
    ws["A1"] = "كشف حساب المقاول — %s" % data["sub"]["Name"]
    ws["A1"].font, ws["A1"].fill = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF"), navy
    ws["A1"].alignment = Alignment(horizontal="right", vertical="center", indent=1, readingOrder=2)
    ws.merge_cells("A2:F2")
    ws["A2"] = "%s  ·  صدر في %s" % (f["label"], datetime.now().strftime("%Y-%m-%d %H:%M"))
    ws["A2"].font, ws["A2"].fill = Font(name="Segoe UI", size=10, color="CBD5E1"), navy
    ws["A2"].alignment = Alignment(horizontal="right", vertical="center", indent=1, readingOrder=2)
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 30, 20
    for j, h in enumerate(("مسلسل", "التاريخ", "الوصف", "مدين", "دائن", "الرصيد"), 1):
        c = ws.cell(4, j, h)
        c.font, c.fill = Font(name="Segoe UI", bold=True, color="FFFFFF"), PatternFill("solid", start_color="0E5A5F")
        c.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)
    row = 5
    if f["date_from"]:
        ws.cell(row, 3, "رصيد سابق").font = Font(name="Segoe UI", bold=True)
        ws.cell(row, 6, data["opening"]).number_format = "#,##0.00"
        row += 1
    for ln in data["lines"]:
        vals = [ln["seq"], ln["date"].strftime("%Y-%m-%d") if ln["date"].year > 1 else "", ln["desc"], ln["debit"] or None, ln["credit"] or None, ln["balance"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row, j, v)
            c.font, c.border = Font(name="Segoe UI", size=10), Border(bottom=line)
            c.alignment = Alignment(horizontal="right" if j == 3 else "center", vertical="center", wrap_text=(j == 3), readingOrder=2)
            if j >= 4:
                c.number_format = "#,##0.00"
        row += 1
    ws.cell(row, 3, "الإجمالي").font = Font(name="Segoe UI", bold=True)
    for j, v in ((4, data["debit"]), (5, data["credit"]), (6, data["balance"])):
        c = ws.cell(row, j, v)
        c.font, c.number_format, c.alignment = Font(name="Segoe UI", bold=True), "#,##0.00", Alignment(horizontal="center")
    ws.cell(row + 1, 3, "الرصيد: %s" % data["note"]).font = Font(name="Segoe UI", bold=True, color="8A2A22" if data["balance"] < 0 else "1A5E3C")
    for j, w in enumerate((8, 13, 70, 16, 16, 16), 1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = "A5"
    ws.page_setup.orientation, ws.page_setup.paperSize = "portrait", 9
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="statement_%s.xlsx" % sub_id, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# --------------------------------------------------------------------------- the register of payments in the archive
@route("/archive/payments")
@auth.login_required
def payments_archive():
    """«الأرشيف ← الدفعات»: every executed payment and movement with its folder and files, to find any of them at once."""
    sub = request.args.get("sub", type=int)
    year = request.args.get("year", type=int)
    rows, subs, years = ls.register(sub, year)
    return render_template("archive_payments.html", rows=rows, subs=subs, years=years, f={"sub": sub or "", "year": year or ""})
