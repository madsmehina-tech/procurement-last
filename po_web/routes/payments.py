# -*- coding: utf-8 -*-
"""Payments tab: every actual transfer to subcontractors, and each contractor's balance (due / paid / remaining)."""
import io
from datetime import datetime

from flask import render_template, request, send_file

import auth
import contracts_service as cs
import db
from routes import route


def _rows():
    a = request.args
    f_sub, f_type, f_cycle = a.get("sub", type=int), a.get("type", ""), a.get("cycle", "").strip()
    subs, contracts, certs, pays = cs.load()
    by_cert = {g["CertificateID"]: g for g in certs}
    cycle_no = {x["CycleID"]: x["CycleNo"] for x in db.query("SELECT [CycleID], [CycleNo] FROM [PaymentCycles]")}
    rows = []
    for p in pays:
        g = by_cert.get(p["CertificateID"]) or {}
        sid = p.get("SubcontractorID") or g.get("SubcontractorID")
        no = cycle_no.get(p.get("CycleID"))
        if f_sub and sid != f_sub:
            continue
        if f_type and f_type != "cert":                     # the only payment type in this company is the certificate
            continue
        if f_cycle and str(no or "") != f_cycle:
            continue
        rows.append({"PaidAt": p.get("PaidAt"), "SubName": (subs.get(sid) or {}).get("Name", ""), "SubcontractorID": sid, "Project": g.get("Project", ""),
                     "SiteName": g.get("SiteName", ""), "CertificateNo": g.get("CertificateNo"), "CertificateID": p["CertificateID"], "CycleID": p.get("CycleID"),
                     "CycleNo": no, "Amount": float(p.get("Amount") or 0), "BankRef": p.get("BankRef") or "", "DocID": p.get("DocID"), "Type": "مستخلص"})
    rows.sort(key=lambda r: (r["PaidAt"] or datetime.min, r["CycleNo"] or 0), reverse=True)
    balances = []
    for sid, s in subs.items():
        f = cs.sub_figures(sid, contracts, certs)
        if f["due"] or f["paid"]:
            balances.append(dict(f, SubName=s["Name"], SubcontractorID=sid))
    balances.sort(key=lambda b: -b["remaining"])
    return rows, balances, subs, dict(sub=f_sub or "", type=f_type, cycle=f_cycle)


@route("/payments")
@auth.login_required
def payments_list():
    rows, balances, subs, f = _rows()
    return render_template("payments.html", rows=rows, balances=balances, f=f, total=sum(r["Amount"] for r in rows),
                           subs=sorted(subs.values(), key=lambda s: s["Name"]))


@route("/payments/export.xlsx")
@auth.login_required
def payments_export():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    rows, balances, _subs, _f = _rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "الدفعات"
    ws.sheet_view.rightToLeft = True
    head = ["التاريخ", "المقاول", "النوع", "المشروع", "الموقع", "المستخلص", "الدورة", "المبلغ", "مرجع التحويل"]
    ws.append(head)
    for r in rows:
        ws.append([r["PaidAt"].strftime("%Y-%m-%d") if r["PaidAt"] else "", r["SubName"], r["Type"], r["Project"], r["SiteName"],
                   r["CertificateNo"], r["CycleNo"] or "", r["Amount"], r["BankRef"]])
    wb2 = wb.create_sheet("أرصدة المقاولين")
    wb2.sheet_view.rightToLeft = True
    wb2.append(["المقاول", "المستحق", "المصروف", "المتبقي"])
    for b in balances:
        wb2.append([b["SubName"], b["due"], b["paid"], b["remaining"]])
    for sheet, widths in ((ws, (12, 32, 10, 24, 18, 10, 8, 14, 20)), (wb2, (32, 16, 16, 16))):
        for i, w in enumerate(widths):
            sheet.column_dimensions[chr(65 + i)].width = w
        for c in sheet[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), PatternFill("solid", start_color="1E293B"), Alignment(horizontal="center")
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "#,##0.00"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="payments_%s.xlsx" % datetime.now().strftime("%Y%m%d"),
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
