# -*- coding: utf-8 -*-
"""Archive tab: every certificate is a transaction; it is classified by itself (contractor -> project -> contract -> certificate)
and is closed when its FIVE papers are complete (closing it also takes it out of the payment cycles). List on the right, the selected transaction's
five papers on the left - the approved reference design."""
import io
from datetime import datetime

from flask import flash, redirect, render_template, request, send_file, url_for

import archive
import audit
import auth
import contracts_service as cs
import db
import ledger_service
from finance import ST_IN_CYCLE, ST_REJECTED
from routes import route
from textnorm import contains

# (type stored in Documents, title, hint) - in the order the company files them
PAPERS = [("نموذج طلب الصرف", "نموذج طلب الصرف", "من الإدارة المالية"), ("المستخلص المعتمد", "المستخلص", "موقّع من المقاول ومدير المشروع"),
          ("الاتفاقية", "الاتفاقية", "عرض سعر / عقد / أمر شراء"), ("الفاتورة الضريبية", "الفاتورة الضريبية", "تُطلب من المقاول قبل التحويل"),
          ("سند الصرف", "سند التحويل البنكي", "بعد إتمام التحويل")]


def transactions():
    subs, contracts, certs, _p = cs.load()
    docs = {}
    for d in db.query("SELECT [DocID], [EntityKey], [DocType], [IsCurrent], [OriginalName] FROM [Documents] WHERE [EntityType] = ?", ("CERT",)):
        if d["IsCurrent"]:
            docs.setdefault(str(d["EntityKey"]), {}).setdefault(d["DocType"], d)
    rows = []
    for g in certs:
        if g["Status"] == ST_REJECTED:
            continue
        mine = docs.get(str(g["CertificateID"]), {})
        contract = contracts.get(g.get("ContractID")) or {}
        papers, count = [], 0
        for key, title, hint in PAPERS:
            d, src = mine.get(key), ""
            if not d and key == "الاتفاقية" and contract.get("ContractFile"):
                d, src = {"DocID": None, "contract": g["ContractID"]}, "من ملف العقد الموقّع"
            papers.append({"key": key, "title": title, "hint": hint, "doc": d, "have": bool(d), "source": src})
            count += 1 if d else 0
        closed = bool(g.get("ClosedAt"))
        complete = count == len(PAPERS)
        rows.append(dict(g, papers=papers, count=count, complete=complete, closed=closed, state="closed" if closed else ("ready" if complete else "open"),
                         can_close=complete and not closed))
    rows.sort(key=lambda r: (r["closed"], -(r.get("CreatedAt") or datetime.min).timestamp() if r.get("CreatedAt") else 0))
    return rows


def _filtered():
    a = request.args
    q, tab = a.get("q", "").strip(), a.get("tab", "all")
    allrows = transactions()
    rows = [r for r in allrows if (tab == "all" or (tab == "closed") == r["closed"]) and
            (not q or any(contains(str(x or ""), q) for x in (r["SubcontractorName"], r["Project"], r.get("SiteName"), r["CertificateNo"], r.get("ContractNo"))))]
    counts = {"all": len(allrows), "open": sum(1 for r in allrows if not r["closed"]), "closed": sum(1 for r in allrows if r["closed"])}
    return allrows, rows, q, tab, counts


@route("/archive")
@auth.login_required
def archive_list():
    _all, rows, q, tab, counts = _filtered()
    sel_id = request.args.get("cert", type=int)
    sel = next((r for r in rows if r["CertificateID"] == sel_id), None) or (rows[0] if rows else None)
    folder = archive.ensure_cert_folder(sel["CertificateID"]) if sel else None
    return render_template("archive_list.html", rows=rows, sel=sel, q=q, tab=tab, counts=counts, n_papers=len(PAPERS), folder=folder, network=archive.network_path(folder) if folder else None)


@route("/archive/export.xlsx")
@auth.login_required
def archive_export():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    _all, rows, _q, _tab, _c = _filtered()
    wb = Workbook()
    ws = wb.active
    ws.title = "أرشيف المعاملات"
    ws.sheet_view.rightToLeft = True
    ws.append(["المقاول", "المشروع", "الموقع", "العقد", "المستخلص", "المصروف"] + [t for _k, t, _h in PAPERS] + ["الحالة"])
    for r in rows:
        ws.append([r["SubcontractorName"], r["Project"], r.get("SiteName") or "", r.get("ContractNo") or "", r["CertificateNo"], r["Paid"]]
                  + ["✓" if p["have"] else "✗" for p in r["papers"]] + [{"closed": "مغلقة", "ready": "جاهزة للإقفال", "open": "مفتوحة (%d/%d)" % (r["count"], len(PAPERS))}[r["state"]]])
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), PatternFill("solid", start_color="14323A"), Alignment(horizontal="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(horizontal="center")
            if isinstance(c.value, float):
                c.number_format = "#,##0.00"
    for i, w in enumerate((28, 24, 18, 16, 10, 14, 18, 14, 18, 16, 18, 16), 1):
        ws.column_dimensions[chr(64 + i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="archive_%s.xlsx" % datetime.now().strftime("%Y%m%d"),
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@route("/archive/<int:cert_id>")
@auth.login_required
def archive_detail(cert_id):
    return redirect(url_for("archive_list", cert=cert_id))


@route("/archive/<int:cert_id>/close", methods=["POST"])
@auth.login_required
def archive_close(cert_id):
    row = next((r for r in transactions() if r["CertificateID"] == cert_id), None)
    if not row:
        flash("لم يتم العثور على المعاملة.")
    elif row["closed"]:
        flash("المعاملة مُقفلة بالفعل.")
    elif not row["complete"]:
        flash("لا تُقفل المعاملة قبل اكتمال أوراقها الخمس (ينقص %d)." % (len(PAPERS) - row["count"]))
    elif row["Status"] == ST_IN_CYCLE:
        flash("هذا المستخلص داخل دورة صرف مفتوحة؛ أكمل الدورة أو احذفها ثم أقفل المعاملة.")
    elif (row.get("Remaining") or 0) > 0.005 and request.form.get("unpaid_ok") != "1":
        # closing is always the owner's free choice, but never by accident: an unpaid balance must be acknowledged
        flash("على هذه المعاملة متبقٍ لم يُصرف (%s ريال). للإقفال مع ذلك أكّد الإقفال من نافذة التأكيد." % "{:,.2f}".format(row["Remaining"]))
    else:
        who, unpaid = (auth.current_user() or {}).get("Username", ""), round(row.get("Remaining") or 0, 2)
        with db.transaction() as tx:                       # the closing and its settlement line are saved together - or neither
            tx.execute("UPDATE [Certificates] SET [ClosedAt] = ?, [ClosedBy] = ? WHERE [CertificateID] = ?", (datetime.now(), who, cert_id))
            if unpaid > 0.005:
                ledger_service.settle_on_close(tx, row, unpaid, who)
        if unpaid > 0.005:
            audit.log("archive_close_unpaid", "مستخلص", cert_id, "أُقفلت المعاملة وسُوّي المتبقي %s ريال (أُقرّ الإقفال)" % "{:,.2f}".format(unpaid))
            flash("أُقفلت المعاملة نهائيًا، وسُجّل المتبقي (%s ريال) «تسوية عند الإقفال» في كشف حساب المقاول فصار المتبقي صفرًا. تُلغى التسوية بإعادة فتح المعاملة." % "{:,.2f}".format(unpaid))
        else:
            flash("أُقفلت المعاملة نهائيًا. أوراقها الخمس محفوظة ولا تُعدَّل إلا بإعادة الفتح من صلاحية أعلى.")
    return redirect(url_for("archive_list", cert=cert_id))


@route("/archive/<int:cert_id>/reopen", methods=["POST"])
@auth.login_required
def archive_reopen(cert_id):
    if db.query("SELECT [CertificateID] FROM [Certificates] WHERE [CertificateID] = ?", (cert_id,)):
        with db.transaction() as tx:
            tx.execute("UPDATE [Certificates] SET [ClosedAt] = NULL, [ClosedBy] = NULL WHERE [CertificateID] = ?", (cert_id,))
            back = ledger_service.remove_settlement(tx, cert_id)
        flash("أُعيد فتح المعاملة للتعديل." + (" وأُلغيت «تسوية عند الإقفال» (%s ريال): عاد المتبقي." % "{:,.2f}".format(back) if back else ""))
    return redirect(url_for("archive_list", cert=cert_id))
