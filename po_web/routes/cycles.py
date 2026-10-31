# -*- coding: utf-8 -*-
"""Payment cycle: the current cycle page (draft / decisions), the statement export screen, the page after approval
(payment orders + transfer follow-up). Built on the approved reference design."""
import io
import logging
from datetime import datetime

from flask import abort, flash, redirect, render_template, request, send_file, url_for

import archive
import auth
import money_guards as mg
import cycle_files
import payment_files
import cycles_service as cy
from routes import route
from routes.documents import save_entity_document

log = logging.getLogger("po.error")


def _user():
    return (auth.current_user() or {}).get("Username", "")


def _num(name):
    return mg.parse_money(request.form.get(name))


def _date(name):
    try:
        return datetime.strptime(request.form.get(name, ""), "%Y-%m-%d")
    except ValueError:
        return datetime.now()


def _get(cycle_id):
    data = cy.load(cycle_id)
    if not data:
        abort(404)
    return data


def _fail(e, cycle_id, endpoint="cycle_detail"):
    """A broken rule is told plainly; anything else is logged and shown with the database's words (nothing was changed)."""
    if isinstance(e, cy.CycleError):
        flash(str(e))
    else:
        log.exception("cycle change failed")
        flash("تعذّر تنفيذ الإجراء ولم يتغير أي شيء. السبب: %s" % " ".join(str(e).split())[:200])
    return redirect(url_for(endpoint, cycle_id=cycle_id))


# --------------------------------------------------------------------------- the cycle tab
@route("/cycles")
@auth.login_required
def cycles_list():
    """«دورة الصرف» opens the CURRENT cycle (the one still being worked on); with none open it offers «إنشاء دورة صرف»."""
    cur = cy.current_cycle()
    if cur:
        return redirect(url_for("cycle_detail", cycle_id=cur["CycleID"]))
    return render_template("cycles_list.html", rows=cy.list_cycles(), ready=cy.candidates(), badge=cy.CYCLE_BADGE, today_iso=datetime.now().strftime("%Y-%m-%d"), history_only=False)


@route("/cycles/all")
@auth.login_required
def cycles_history():
    return render_template("cycles_list.html", rows=cy.list_cycles(), ready=[], badge=cy.CYCLE_BADGE, today_iso=datetime.now().strftime("%Y-%m-%d"), history_only=True)


@route("/cycles/new", methods=["GET", "POST"])
@auth.login_required
def cycle_new():
    if request.method == "GET":
        return redirect(url_for("cycles_list"))
    try:
        cid = cy.create(request.form.get("title", "").strip(), _date("cycle_date"), _user())
    except cy.CycleError as e:
        flash(str(e))
        return redirect(url_for("cycles_list"))
    except Exception as e:                                   # noqa
        log.exception("cycle create failed")
        flash("تعذّر إنشاء الدورة ولم يتغير أي شيء. السبب: %s" % " ".join(str(e).split())[:200])
        return redirect(url_for("cycles_list"))
    archive.ensure_cycle_folder(cid)
    flash("أُنشئت الدورة من كل المستخلصات المفتوحة. راجع الكشف ثم اضغط «الانتقال للمرحلة التالية».")
    return redirect(url_for("cycle_detail", cycle_id=cid))


# --------------------------------------------------------------------------- the cycle page (the three stages)
@route("/cycles/<int:cycle_id>")
@auth.login_required
def cycle_detail(cycle_id):
    data = _get(cycle_id)
    c = data["cycle"]
    orders = [gp for gp in data["groups"] if gp["payable"] > cy.EPS]
    docs = archive.panel("CYCLE", cycle_id, archive.CYCLE_DOC_TYPES)
    folder = archive.ensure_cycle_folder(cycle_id)
    can_delete = not any(l["Pays"] for l in data["lines"]) and not any(l["Exec"] == cy.EX_DONE for l in data["lines"])
    payinfo = {}
    for gp in data["groups"]:
        no = cy.group_pay_no(cycle_id, gp["sub_id"]) if gp["exec"] == cy.EX_DONE else None
        if no:
            payinfo[gp["sub_id"]] = {"no": no, "docs": archive.list_documents("PAYMENT", "%s:%s" % (gp["sub_id"], no)), "folder": archive.payment_folder_for(gp["sub_id"], no)}
    return render_template("cycle_detail.html", d=data, badge=cy.CYCLE_BADGE, ST=cy, STAGES=cy.STAGES, D_FULL=cy.D_FULL, D_PART=cy.D_PART, D_DEFER=cy.D_DEFER,
                           DECISION_LABEL=cy.DECISION_LABEL, EX_DONE=cy.EX_DONE, EX_NOT=cy.EX_NOT, today_iso=datetime.now().strftime("%Y-%m-%d"), orders=orders,
                           docs=docs, folder=folder, network=archive.network_path(folder), can_delete=can_delete, payinfo=payinfo,
                           deciding=c["Status"] == cy.ST_WAITING, noting=c["Status"] == cy.ST_DRAFT)


@route("/cycles/<int:cycle_id>/decisions", methods=["POST"])
@auth.login_required
def cycle_decisions(cycle_id):
    """Saves what is on the page. intent = save (stay) / next (stage 1 -> 2) / finalize (stage 2 -> 3)."""
    data = _get(cycle_id)
    entries = {}
    intent = request.form.get("intent", "save")
    try:
        for l in data["lines"]:
            did = l["DecisionID"]
            partial = mg.parse_money(request.form.get("partial_%d" % did), label="مبلغ الصرف الجزئي")
            entries[did] = (request.form.get("decision_%d" % did, ""), partial, request.form.get("note_%d" % did, ""))
        n = cy.save_lines(cycle_id, entries, _user())
        if intent.startswith("distribute:"):
            sub_id = int(intent.split(":", 1)[1])
            raw = (request.form.get("dist_%d" % sub_id) or "").strip()
            if not raw:
                raise cy.CycleError("اكتب «مبلغ الصرف المعتمد» لهذا المقاول أولًا.")
            plan = cy.distribute(cycle_id, sub_id, mg.parse_money(raw, label="مبلغ الصرف المعتمد"), _user())
            full, part, defer = (sum(1 for _d, d, _a in plan if d == x) for x in (cy.D_FULL, cy.D_PART, cy.D_DEFER))
            flash("وُزّع المبلغ على مستخلصات المقاول (الأقدم أولًا): %d بالكامل، %d جزئيًا، %d مؤجل. يمكنك تعديل أي قرار أو مبلغ يدويًا ثم «حفظ»." % (full, part, defer))
        elif intent == "next":
            cy.to_decision_stage(cycle_id, _user())
            doc, err = cycle_files.save_sheet(cycle_id, "blank", None, _user())
            flash("انتقلت الدورة لمرحلة قرار الإدارة. اطبع الكشف واعرضه على الإدارة، ثم سجّل القرار والمبلغ لكل مستخلص%s." % (
                " (وحُفظ PDF الكشف في أرشيف الدورة)" if doc else " (لم يُنشأ PDF: %s)" % err))
        elif intent == "finalize":
            cy.to_execution_stage(cycle_id, _user())
            doc, err = cycle_files.save_sheet(cycle_id, "values", None, _user())
            flash("انتقلت الدورة للمرحلة النهائية: القرارات مقفلة ويتبقى تحديد «تم التحويل / لم يتم التنفيذ» لكل مستخلص%s." % (
                " (وحُفظ PDF الكشف في أرشيف الدورة)" if doc else " (لم يُنشأ PDF: %s)" % err))
        else:
            flash("تم الحفظ%s." % (" (%d من %d عليها قرار)" % (n, len(data["lines"])) if data["cycle"]["Status"] == cy.ST_WAITING else ""))
    except mg.MoneyError as e:
        flash(str(e))
    except ValueError:
        flash("مبلغ الصرف يجب أن يكون رقمًا.")
    except Exception as e:                                   # noqa
        return _fail(e, cycle_id)
    return redirect(url_for("cycle_detail", cycle_id=cycle_id))


@route("/cycles/<int:cycle_id>/refresh", methods=["POST"])
@auth.login_required
def cycle_refresh(cycle_id):
    try:
        cy.refresh(cycle_id, _user())
        flash("أُعيد احتساب قائمة الدورة من المستخلصات المفتوحة الحالية (الملاحظات المكتوبة بقيت).")
    except Exception as e:                                   # noqa
        return _fail(e, cycle_id)
    return redirect(url_for("cycle_detail", cycle_id=cycle_id))


@route("/cycles/<int:cycle_id>/close", methods=["POST"])
@auth.login_required
def cycle_close(cycle_id):
    """Closes the cycle once every approved line has an outcome."""
    try:
        cy.close_cycle(cycle_id, _user())
        flash("تم إغلاق الدورة.")
    except Exception as e:                                   # noqa
        return _fail(e, cycle_id)
    return redirect(url_for("cycle_detail", cycle_id=cycle_id))


@route("/cycles/<int:cycle_id>/reopen", methods=["POST"])
@auth.login_required
def cycle_reopen(cycle_id):
    """One stage back, at any time (a closed cycle reopens, execution -> decisions, decisions -> creation). Logged."""
    try:
        cy.go_back(cycle_id, _user())
        flash("رجعت الدورة مرحلة واحدة. كل تغيير يُسجَّل في سجل الدورة.")
    except Exception as e:                                   # noqa
        return _fail(e, cycle_id)
    return redirect(url_for("cycle_detail", cycle_id=cycle_id))


@route("/cycles/<int:cycle_id>/delete", methods=["POST"])
@auth.login_required
def cycle_delete(cycle_id):
    """Only while no transfer has been recorded in the cycle."""
    try:
        no = cy.delete_cycle(cycle_id, _user())
        flash("حُذفت دورة الصرف رقم %s وعادت مستخلصاتها لقائمة المستخلصات المفتوحة (سجل الدورة محفوظ)." % no)
    except Exception as e:                                   # noqa
        return _fail(e, cycle_id)
    return redirect(url_for("cycles_list"))


# --------------------------------------------------------------------------- the statement: print, PDF, Excel, signed copy
@route("/cycles/<int:cycle_id>/export")
@auth.login_required
def cycle_export(cycle_id):
    data = _get(cycle_id)
    mode, cols = cycle_files.parse_options(request.args, cycle_files.default_mode_for(data["cycle"]))
    ctx = cycle_files.sheet_context(data, mode, cols)
    return render_template("cycle_export.html", COLS=cycle_files.COLS, chosen=cols, ST=cy, **ctx)


@route("/cycles/<int:cycle_id>/print")
@auth.login_required
def cycle_print(cycle_id):
    data = _get(cycle_id)
    mode, cols = cycle_files.parse_options(request.args, cycle_files.default_mode_for(data["cycle"]))
    return cycle_files.sheet_html(data, mode, cols)


@route("/cycles/<int:cycle_id>/sheet.pdf")
@auth.login_required
def cycle_pdf_download(cycle_id):
    """The statement as a PDF file to download (A4 landscape). Without Edge / Chrome on the PC the print page opens instead."""
    data = _get(cycle_id)
    mode, cols = cycle_files.parse_options(request.args, cycle_files.default_mode_for(data["cycle"]))
    try:
        pdf = cycle_files._pdf_bytes(cycle_files.sheet_html(data, mode, cols, pdf=True))
    except cycle_files.pdfgen.PdfError as e:
        flash("تعذّر إنشاء PDF تلقائيًا (%s) - استخدم «طباعة الكشف» ثم «حفظ كـ PDF» من المتصفح." % e)
        return redirect(url_for("cycle_export", cycle_id=cycle_id, mode=mode, cols=",".join(cols)))
    return send_file(io.BytesIO(pdf), as_attachment=True, download_name="cycle_%s.pdf" % data["cycle"]["CycleNo"], mimetype="application/pdf")


@route("/cycles/<int:cycle_id>/pdf", methods=["POST"])
@auth.login_required
def cycle_pdf(cycle_id):
    mode, cols = cycle_files.parse_options(request.form)
    doc, err = cycle_files.save_sheet(cycle_id, mode, cols, _user())
    flash("حُفظ PDF الكشف في أرشيف الدورة (الكشوف)." if doc else "لم يُنشأ PDF: %s" % err)
    return redirect(url_for("cycle_export", cycle_id=cycle_id, mode=mode, cols=",".join(cols)))


@route("/cycles/<int:cycle_id>/signed", methods=["POST"])
@auth.login_required
def cycle_signed(cycle_id):
    """«رفع مستندات الدورة المعتمدة»: the signed copy of the cycle and any attachment, into the cycle's archive folder."""
    _get(cycle_id)
    doc_type = request.form.get("doc_type") if request.form.get("doc_type") in ("الكشف الموقّع", "مرفق إضافي") else "الكشف الموقّع"
    doc = save_entity_document("CYCLE", cycle_id, doc_type, request.files.get("file"))
    if doc:
        flash("حُفظ «%s» في أرشيف الدورة." % doc_type)
    return redirect(url_for("cycle_detail", cycle_id=cycle_id))


@route("/cycles/<int:cycle_id>/export.xlsx")
@auth.login_required
def cycle_xlsx(cycle_id):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    data = _get(cycle_id)
    mode, cols = cycle_files.parse_options(request.args, cycle_files.default_mode_for(data["cycle"]))
    c, t = data["cycle"], data["totals"]
    blank = mode == "blank"
    wb = Workbook()
    ws = wb.active
    ws.title = "كشف الدورة"
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    heads = [h for k, h, _w in cycle_files.COLS if k in cols]
    navy, line = PatternFill("solid", start_color="14323A"), Side(style="thin", color="E2E8F0")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(heads))
    ws["A1"] = "كشف دورة الصرف رقم %s — شركة الهضب للتجارة والمقاولات" % c["CycleNo"]
    ws["A1"].font, ws["A1"].fill = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF"), navy
    ws["A1"].alignment = Alignment(horizontal="right", vertical="center", indent=1, readingOrder=2)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(heads))
    ws["A2"] = "%s  ·  %s  ·  التاريخ: %s  ·  إجمالي المتبقي %s  ·  مبلغ الصرف %s ريال" % (c.get("Title") or "", cy.STAGE_NAME.get(c["Status"], c["Status"]),
                                                                                       (c.get("CycleDate") or c["CreatedAt"]).strftime("%Y-%m-%d"),
                                                                                       "{:,.2f}".format(t["all"]), "{:,.2f}".format(t["approved"]))
    ws["A2"].font, ws["A2"].fill = Font(name="Segoe UI", size=10, color="CBD5E1"), navy
    ws["A2"].alignment = Alignment(horizontal="right", vertical="center", indent=1, readingOrder=2)
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 30, 20
    for j, h in enumerate(heads, 1):
        cell = ws.cell(4, j, h)
        cell.font, cell.fill = Font(name="Segoe UI", bold=True, color="FFFFFF"), PatternFill("solid", start_color="0E5A5F")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
    money = {"value", "paid_before", "remaining", "approved", "total", "sub_approved"}
    first_of = {gp["lines"][0]["CertificateID"]: gp for gp in data["groups"]}        # «الإجمالي» / «مبلغ الصرف المعتمد»: one value per contractor, on its first row
    starts = {}
    for i, l in enumerate(data["lines"], 1):
        decided = bool(l["Decision"]) and not blank
        val = {"seq": i, "sub": l["SubName"], "ref": l["CertNo"], "project": l["Project"], "site": l["Site"], "value": l["Total"], "paid_before": l["PaidBefore"],
               "remaining": l["Requested"], "total": None, "sub_approved": None, "status": l["StatusLabel"], "decision": l["DecisionLabel"] if decided else "", "approved": l["Approved"] if decided else None,
               "notes": "" if blank else l["Note"]}
        gp = first_of.get(l["CertificateID"])
        if gp:
            val["total"] = gp["total"]
            val["sub_approved"] = None if blank or not gp["decided"] else gp["approved"]
            starts[gp["sub_id"]] = (4 + i, len(gp["lines"]))
        for j, k in enumerate([k for k, _h, _w in cycle_files.COLS if k in cols], 1):
            cell = ws.cell(4 + i, j, val[k])
            cell.font, cell.border = Font(name="Segoe UI", size=10), Border(bottom=line)
            cell.alignment = Alignment(horizontal="right" if k in ("sub", "project", "site", "notes") else "center", readingOrder=2, wrap_text=k in ("sub", "project", "site", "notes"))
            if k in money:
                cell.number_format = "#,##0.00"
    last = 4 + len(data["lines"])
    keys = [k for k, _h, _w in cycle_files.COLS if k in cols]
    for sub_id, (row, n) in starts.items():                  # the contractor's two merged cells
        for k in cycle_files.GROUP_COLS:
            if k in keys and n > 1:
                col = keys.index(k) + 1
                ws.merge_cells(start_row=row, start_column=col, end_row=row + n - 1, end_column=col)
                ws.cell(row, col).alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)
                ws.cell(row, col).font = Font(name="Segoe UI", size=10, bold=True)
    ws.cell(last + 1, 2, "الإجمالي").font = Font(name="Segoe UI", bold=True)
    for k, v in (("value", sum(l["Total"] for l in data["lines"])), ("paid_before", sum(l["PaidBefore"] for l in data["lines"])), ("remaining", t["all"]),
                 ("approved", None if blank else t["approved"])):
        if k in keys and v is not None:
            cell = ws.cell(last + 1, keys.index(k) + 1, round(v, 2))
            cell.font, cell.number_format = Font(name="Segoe UI", bold=True), "#,##0.00"
    widths = {"seq": 8, "sub": 30, "ref": 10, "project": 24, "site": 18, "value": 16, "paid_before": 16, "remaining": 16, "total": 16, "sub_approved": 18, "status": 20, "decision": 14, "approved": 16, "notes": 30}
    for j, k in enumerate(keys, 1):
        ws.column_dimensions[chr(64 + j)].width = widths[k]
    ws.freeze_panes = "A5"
    ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", 9
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="cycle_%s.xlsx" % c["CycleNo"], mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# --------------------------------------------------------------------------- payment orders (kept as they are - to be discussed before launch)
@route("/cycles/<int:cycle_id>/subs/<int:sub_id>/order")
@auth.login_required
def cycle_order(cycle_id, sub_id):
    html = cycle_files.order_html(_get(cycle_id), sub_id)
    if not html:
        flash("لا مبلغ معتمد للصرف لهذا المقاول.")
        return redirect(url_for("cycle_detail", cycle_id=cycle_id))
    return html


@route("/cycles/<int:cycle_id>/orders")
@auth.login_required
def cycle_orders_print(cycle_id):
    html = cycle_files.all_orders_html(_get(cycle_id))
    if not html:
        flash("لا أوامر صرف في هذه الدورة بعد.")
        return redirect(url_for("cycle_detail", cycle_id=cycle_id))
    return html


@route("/cycles/<int:cycle_id>/subs/<int:sub_id>/order-pdf", methods=["POST"])
@auth.login_required
def cycle_order_pdf(cycle_id, sub_id):
    doc, err = cycle_files.save_order(cycle_id, sub_id, _user())
    flash("حُفظ أمر الصرف (PDF) في «أوامر الصرف» داخل مجلد الدورة." if doc else "لم يُنشأ PDF: %s" % err)
    return redirect(url_for("cycle_detail", cycle_id=cycle_id))


# --------------------------------------------------------------------------- the last stage: ONE transfer per contractor
@route("/cycles/<int:cycle_id>/subs/<int:sub_id>/transfer", methods=["POST"])
@auth.login_required
def cycle_transfer(cycle_id, sub_id):
    """The outcome of the contractor's transfer. «تم التحويل»: every certificate of the contractor gets its payment, the payment gets its
    number and its archive folder, and the files uploaded now (the signed order + the bank transfer picture) are kept in that folder and
    in each certificate paid (as «سند الصرف»). Files can also be added later. «لم يتم التنفيذ»: nothing is paid, nothing is archived."""
    state = request.form.get("state", "")
    files = [f for f in request.files.getlist("proof_file") if f and f.filename]
    if state == cy.EX_DONE and payment_files.bad_file(files):
        flash("صيغة الملف غير مدعومة: %s - لم يُسجَّل التنفيذ." % payment_files.bad_file(files))
        return redirect(url_for("cycle_detail", cycle_id=cycle_id))
    try:
        status = cy.set_group_execution(cycle_id, sub_id, state, _date("pay_date"), request.form.get("bank_ref", "").strip(), request.form.get("note", ""), _user())
    except Exception as e:                                         # noqa
        return _fail(e, cycle_id)
    msg = "سُجّل «%s» للمقاول" % state
    if state == cy.EX_DONE:
        try:
            pay_no, _folder = payment_files.after_execution(cycle_id, sub_id)
            msg += " (دفعة رقم %s)" % pay_no
            if files:
                n = payment_files.store_transfer_files(cycle_id, sub_id, files, request.form.get("note", "").strip(), _user())
                msg += " - رُفع %d ملف إلى أرشيف الدفعة وإلى كل مستخلص" % n
        except Exception as e:                                     # noqa - the transfer is recorded; only the archive step failed
            msg += " - لكن تعذّر حفظ الملفات (%s)؛ يمكن رفعها من «رفع ملفات التحويل»" % " ".join(str(e).split())[:120]
    flash(msg + (" - وبهذا اكتملت الدورة وأُغلقت." if status == cy.ST_DONE else "."))
    return redirect(url_for("cycle_detail", cycle_id=cycle_id))


@route("/cycles/<int:cycle_id>/subs/<int:sub_id>/files", methods=["POST"])
@auth.login_required
def cycle_transfer_files(cycle_id, sub_id):
    """Adds the scan of the executed transfer afterwards (the order is signed and the bank picture arrives after «تم التحويل»)."""
    files = [f for f in request.files.getlist("proof_file") if f and f.filename]
    if not files:
        flash("اختر ملفًا أولًا.")
        return redirect(url_for("cycle_detail", cycle_id=cycle_id))
    try:
        n = payment_files.store_transfer_files(cycle_id, sub_id, files, request.form.get("note", "").strip(), _user())
        flash("رُفع %d ملف إلى أرشيف الدفعة وإلى كل مستخلص دُفع بها." % n)
    except Exception as e:                                         # noqa
        flash(" ".join(str(e).split())[:200])
    return redirect(url_for("cycle_detail", cycle_id=cycle_id))


@route("/cycles/<int:cycle_id>/subs/<int:sub_id>/undo", methods=["POST"])
@auth.login_required
def cycle_transfer_undo(cycle_id, sub_id):
    try:
        cy.clear_group_execution(cycle_id, sub_id, _user())
        flash("أُلغيت حالة التنفيذ لهذا المقاول (ملفات معاملاته باقية في الأرشيف) ويمكنك تحديدها من جديد.")
    except Exception as e:                                         # noqa
        return _fail(e, cycle_id)
    return redirect(url_for("cycle_detail", cycle_id=cycle_id))
