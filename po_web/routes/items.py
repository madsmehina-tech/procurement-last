# -*- coding: utf-8 -*-
"""Items tab: price history of every item bought, built from the purchase orders - with project / site filters and professional exports
(Excel tables for Pivot / Power BI, and a PDF report) of the item x project x site analysis."""
import io
from datetime import datetime

from flask import flash, redirect, render_template, request, send_file, url_for

import auth
import cycle_files
import items_service
from routes import route
from textnorm import contains


def _filtered():
    a = request.args
    q, f_supplier, f_project, f_site = a.get("q", "").strip(), a.get("supplier", "").strip(), a.get("project", "").strip(), a.get("site", "").strip()
    items = items_service.build(f_project, f_site)
    suppliers = sorted({h["supplier"] for i in items_service.build() for h in i["history"]})
    if q:
        items = [i for i in items if contains(i["name"], q)]
    if f_supplier:
        for i in items:                                   # keep only that supplier's purchases; the totals follow
            i["history"] = [h for h in i["history"] if h["supplier"] == f_supplier]
        items = [i for i in items if i["history"]]
        items = _rebuilt(items)
    label = " · ".join(x for x in (("مشروع: " + f_project) if f_project else "", ("موقع: " + f_site) if f_site else "", ("مورد: " + f_supplier) if f_supplier else "",
                                   ("بحث: " + q) if q else "") if x) or "كل المشتريات"
    return items, suppliers, dict(q=q, supplier=f_supplier, project=f_project, site=f_site, label=label)


def _rebuilt(items):
    """Totals of items whose purchases were narrowed down to one supplier."""
    for i in items:
        h = i["history"]
        i["count"], i["total_qty"] = len(h), sum(x["qty"] for x in h)
        i["total_value"] = round(sum(x["total"] for x in h), 2)
        prices = [x["price"] for x in h if x["price"]]
        i["avg_price"] = round(i["total_value"] / i["total_qty"], 2) if i["total_qty"] else 0
        i["last_price"], i["max_price"], i["min_price"] = h[0]["price"], (max(prices) if prices else 0), (min(prices) if prices else 0)
        i["last_supplier"] = i["top_supplier"] = h[0]["supplier"]
    return items


@route("/items")
@auth.login_required
def items_list():
    items, suppliers, f = _filtered()
    sort = request.args.get("sort", "name")
    desc = request.args.get("dir", "asc" if sort == "name" else "desc") == "desc"
    view = request.args.get("view", "items") if request.args.get("view") in ("items", "analysis") else "items"
    items.sort(key=items_service.SORTS.get(sort, items_service.SORTS["name"]), reverse=desc)
    rows = items_service.analysis(items)
    projects, sites_by_project = items_service.filter_options()
    all_sites = sorted({s for v in sites_by_project.values() for s in v}) if not f["project"] else sites_by_project.get(f["project"], [])
    kpi = {"count": len(items), "purchases": sum(i["count"] for i in items), "value": sum(i["total_value"] for i in items),
           "top": max(items, key=lambda i: i["total_value"])["name"] if items else "-", "rows": len(rows)}
    return render_template("items_list.html", items=items, rows=rows, view=view, f=f, q=f["q"], f_supplier=f["supplier"], sort=sort, desc=desc,
                           suppliers=suppliers, kpi=kpi, projects=projects, sites=all_sites)


@route("/items/detail")
@auth.login_required
def item_detail():
    key = request.args.get("k", "")
    item = next((i for i in items_service.build() if i["key"] == key), None)
    if not item:
        flash("لم يتم العثور على الصنف.")
        return redirect(url_for("items_list"))
    return render_template("item_detail.html", item=item, spark=items_service.sparkline(item["history"]))


def _report_html(pdf=False):
    items, _suppliers, f = _filtered()
    rows = items_service.analysis(items)
    return render_template("items_print.html", rows=rows, f=f, now=datetime.now().strftime("%Y-%m-%d %H:%M"),
                           total_qty=sum(r["qty"] for r in rows), total_value=sum(r["value"] for r in rows), n_items=len(items))


@route("/items/print")
@auth.login_required
def items_print():
    return _report_html()


@route("/items/export.pdf")
@auth.login_required
def items_export_pdf():
    try:
        pdf = cycle_files._pdf_bytes(_report_html(pdf=True))
    except cycle_files.pdfgen.PdfError as e:
        flash("تعذّر إنشاء PDF تلقائيًا (%s) - استخدم «طباعة» ثم «حفظ كـ PDF» من المتصفح." % e)
        return redirect(url_for("items_list", view="analysis", **request.args))
    return send_file(io.BytesIO(pdf), as_attachment=True, download_name="items_%s.pdf" % datetime.now().strftime("%Y%m%d"), mimetype="application/pdf")


@route("/items/export.xlsx")
@auth.login_required
def items_export_xlsx():
    """Three sheets, each a real Excel TABLE (Pivot / Power BI read them by name):
    tblAnalysis (item x project x site), tblItems (one row per item), tblPurchases (every purchase line)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    items, _suppliers, f = _filtered()
    rows = items_service.analysis(items)
    wb = Workbook()

    def sheet(ws, title, band, headers, data, widths, money_cols, name):
        ws.title = title
        ws.sheet_view.rightToLeft = True
        ws.sheet_view.showGridLines = False
        n = len(headers)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
        ws["A1"] = "شركة الهضب للتجارة والمقاولات — " + band
        ws["A1"].font, ws["A1"].fill = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF"), PatternFill("solid", start_color="14323A")
        ws["A1"].alignment = Alignment(horizontal="right", vertical="center", indent=1, readingOrder=2)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
        ws["A2"] = "%s  ·  صدر في %s" % (f["label"], datetime.now().strftime("%Y-%m-%d %H:%M"))
        ws["A2"].font, ws["A2"].fill = Font(name="Segoe UI", size=10, color="CBD5E1"), PatternFill("solid", start_color="14323A")
        ws["A2"].alignment = Alignment(horizontal="right", vertical="center", indent=1, readingOrder=2)
        ws.row_dimensions[1].height, ws.row_dimensions[2].height = 28, 18
        for j, h in enumerate(headers, 1):
            ws.cell(4, j, h)
        for i, row in enumerate(data, 5):
            for j, v in enumerate(row, 1):
                c = ws.cell(i, j, v)
                c.font = Font(name="Segoe UI", size=10)
                c.alignment = Alignment(horizontal="center" if j in money_cols or isinstance(v, (int, float)) else "right", vertical="center", readingOrder=2)
                if j in money_cols:
                    c.number_format = "#,##0.00"
        last = 4 + max(len(data), 1)
        if data:
            tbl = Table(displayName=name, ref="A4:%s%d" % (get_column_letter(n), last))
            tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(tbl)
        else:
            for j in range(1, n + 1):
                ws.cell(4, j).font = Font(bold=True)
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "A5"
        ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", 9
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "4:4"
        return last

    ws = wb.active
    last = sheet(ws, "تحليل الأصناف", "تحليل الأصناف (صنف × مشروع × موقع)",
                 ["اسم الصنف", "الوحدة", "المشروع", "الموقع", "إجمالي الكمية", "إجمالي القيمة", "متوسط السعر", "عدد أوامر الشراء"],
                 [[r["item"], r["unit"], r["project"], r["site"], r["qty"], r["value"], r["avg"], r["orders"]] for r in rows], (34, 10, 28, 24, 14, 16, 14, 12), (6, 7), "tblAnalysis")
    ws.cell(last + 2, 4, "الإجمالي").font = Font(name="Segoe UI", bold=True)
    for col, v in ((5, sum(r["qty"] for r in rows)), (6, sum(r["value"] for r in rows))):
        c = ws.cell(last + 2, col, round(v, 2))
        c.font, c.number_format = Font(name="Segoe UI", bold=True), "#,##0.00"
    sheet(wb.create_sheet(), "ملخص الأصناف", "ملخص الأصناف",
          ["اسم الصنف", "الوحدة", "مرات الشراء", "إجمالي الكمية", "إجمالي القيمة", "متوسط السعر", "آخر سعر", "أعلى سعر", "أقل سعر", "أكثر الموردين توريدًا", "آخر مورد"],
          [[i["name"], i["unit"], i["count"], i["total_qty"], i["total_value"], i["avg_price"], i["last_price"], i["max_price"], i["min_price"], i["top_supplier"], i["last_supplier"]]
           for i in sorted(items, key=lambda i: i["name"])], (34, 10, 12, 14, 16, 14, 14, 14, 14, 28, 28), (5, 6, 7, 8, 9), "tblItems")
    purchases = sorted(((h, i) for i in items for h in i["history"]), key=lambda t: (t[0]["date"] or datetime.min, t[0]["po"]), reverse=True)
    sheet(wb.create_sheet(), "تفاصيل الشراء", "كل أسطر الشراء",
          ["التاريخ", "أمر الشراء", "المورد", "اسم الصنف", "الوحدة", "المشروع", "الموقع", "الكمية", "سعر الوحدة", "الإجمالي"],
          [[h["date"].strftime("%Y-%m-%d") if h["date"] else "", h["po"], h["supplier"], i["name"], i["unit"], h["project"], h["site"], h["qty"], h["price"], h["total"]] for h, i in purchases],
          (12, 16, 28, 34, 10, 28, 24, 12, 14, 16), (8, 9, 10), "tblPurchases")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="items_%s.xlsx" % datetime.now().strftime("%Y%m%d"), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
