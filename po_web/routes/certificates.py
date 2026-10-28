# -*- coding: utf-8 -*-
"""
Certificates (المستخلصات). The approved design: the accounting review happens outside the system as the company does
it today; a reviewed certificate is ENTERED here (saved + archived), management decides on it in the payment cycle.
    مُدخَل  ->  ضمن دورة صرف  ->  مصروف        (مرفوض = returned, with a reason; it can be re-opened)
Until the payment-cycle screens exist (next batch), a payment can be recorded directly on a certificate.
"""
import logging
from datetime import datetime

from flask import abort, flash, redirect, render_template, request, url_for

import archive
import auth
import contracts_service as cs
import db
import masterdata
import money_guards as mg
import opening_service
from finance import STATUS_BADGE, ST_ENTERED, ST_IN_CYCLE, ST_PAID, ST_REJECTED, CERT_STATUSES, classify_status, norm_status
from routes import route
from routes.documents import save_entity_document
from textnorm import clean_spaces, contains, norm_key



def _num(name, default=0.0, label="المبلغ"):
    """An amount from the form: finite, not negative, two decimals (raises mg.MoneyError, a ValueError, with a message for the person)."""
    raw = (request.form.get(name) or "").strip()
    return mg.parse_money(raw, label=label) if raw else default


def _date(name):
    try:
        return datetime.strptime(request.form.get(name, ""), "%Y-%m-%d")
    except ValueError:
        return datetime.now()


def _user():
    return (auth.current_user() or {}).get("Username", "")


# --------------------------------------------------------------------------- the certificates tab (follow-up only)
def _filtered():
    a = request.args
    q, f_status, f_sub, f_project = a.get("q", "").strip(), a.get("status", "").strip(), a.get("sub", type=int), a.get("project", "").strip()
    d_from, d_to = a.get("date_from", ""), a.get("date_to", "")
    subs, contracts, certs, _p = cs.load()
    rows = []
    for g in sorted(certs, key=lambda g: (g.get("CreatedAt") or datetime.min, g["CertificateID"]), reverse=True):
        if f_status and g["Status"] != f_status:
            continue
        if f_sub and g["SubcontractorID"] != f_sub:
            continue
        if f_project and g["Project"] != f_project:
            continue
        when = g.get("CreatedAt") or g.get("CertificateDate")
        if d_from and when and when.strftime("%Y-%m-%d") < d_from:
            continue
        if d_to and when and when.strftime("%Y-%m-%d") > d_to:
            continue
        if q and not any(contains(str(x or ""), q) for x in (g["CertificateNo"], g["SubcontractorName"], g["Project"], g.get("SiteName"), g.get("ContractNo"))):
            continue
        rows.append(g)
    return rows, subs, contracts, dict(q=q, status=f_status, sub=f_sub or "", project=f_project, date_from=d_from, date_to=d_to)


@route("/certificates")
@auth.login_required
def certificates_list():
    rows, subs, contracts, f = _filtered()
    by_status = {s: {"n": 0, "total": 0.0} for s in CERT_STATUSES}
    for g in rows:
        slot = by_status.setdefault(g["Status"], {"n": 0, "total": 0.0})        # never a KeyError on an odd value
        slot["n"] += 1
        slot["total"] += g["Total"]
    return render_template("certificates_list.html", rows=rows, by_status=by_status, statuses=CERT_STATUSES, badge=STATUS_BADGE, f=f,
                           subs=sorted(subs.values(), key=lambda s: s["Name"]), projects=sorted({g.get("Project") for g in cs.load()[2] if g.get("Project")}))


@route("/certificates/export.xlsx")
@auth.login_required
def certificates_export():
    import io
    from flask import send_file
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    rows, _subs, _contracts, _f = _filtered()
    wb = Workbook()
    ws = wb.active
    ws.title = "المستخلصات"
    ws.sheet_view.rightToLeft = True
    ws.append(["مسلسل", "رقم المستخلص", "اسم المقاول", "المشروع", "الموقع", "تاريخ الإدخال للنظام", "قيمة المستخلص شامل الضريبة", "خصومات المواد",
               "صافي المستحق", "المصروف", "رقم دورة الصرف", "تاريخ الصرف", "المتبقي", "الحالة"])
    for i, g in enumerate(rows, 1):
        ws.append([i, g["CertificateNo"], g["SubcontractorName"], g["Project"], g.get("SiteName") or "", g["CreatedAt"].strftime("%Y-%m-%d") if g.get("CreatedAt") else "",
                   g["Value"], g.get("SuppliedMaterials") or 0, g["Total"], g["Paid"], g.get("CycleNo") or "", g["LastPayDate"].strftime("%Y-%m-%d") if g.get("LastPayDate") else "",
                   g["Remaining"], g["Status"]])
    for c_ in ws[1]:
        c_.font, c_.fill, c_.alignment = Font(bold=True, color="FFFFFF"), PatternFill("solid", start_color="1E293B"), Alignment(horizontal="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c_ in row:
            c_.alignment = Alignment(horizontal="center")
            if isinstance(c_.value, float):
                c_.number_format = "#,##0.00"
    for i, w in enumerate((8, 12, 30, 24, 18, 14, 18, 14, 16, 16, 12, 14, 16, 14), 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="certificates_%s.xlsx" % datetime.now().strftime("%Y%m%d"),
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# --------------------------------------------------------------------------- add / edit / delete: from the contractor's page ONLY
def _back_to_sub(sub_id, **kw):
    return redirect(url_for("subcontractor_detail", sub_id=sub_id, **kw) + "#certificates")


def _read_cert_form(sub_id, existing=None):
    """Reads and checks the certificate form. Returns (values, errors).
    A certificate stands on its own: contractor + project + site (names typed or picked from the unified lists); the contract is an
    optional link. A new project or site is saved to the lists when the certificate is saved."""
    f, errors = request.form, []
    v = {"contract_id": None}
    if existing:
        v["project"], v["site"] = existing.get("Project") or "", existing.get("SiteName") or ""
    else:
        v["project"], v["site"] = clean_spaces(f.get("project", "")), clean_spaces(f.get("site", ""))
        if not v["project"]:
            errors.append("اسم المشروع مطلوب.")
        if not v["site"]:
            errors.append("اسم الموقع مطلوب.")
    raw = (f.get("contract_id") or "").strip()                 # optional
    if raw:
        mine = {str(c["ContractID"]) for c in db.query("SELECT [ContractID] FROM [Contracts] WHERE [SubcontractorID] = ?", (sub_id,))}
        if raw in mine:
            v["contract_id"] = int(raw)
        else:
            errors.append("العقد المختار ليس من عقود هذا المقاول.")
    elif existing and "contract_id" not in f:
        v["contract_id"] = existing.get("ContractID")          # a form that does not carry the field keeps the link as it is
    try:
        v["value"], v["materials"], v["penalties"] = _num("total_incl_vat", label="إجمالي المستخلص"), _num("supplied_materials", label="مواد الشركة"), _num("penalties", label="الغرامات")
    except mg.MoneyError as e:
        v["value"] = v["materials"] = v["penalties"] = 0.0
        errors.append(str(e))
    if v["value"] <= 0:
        errors.append("إجمالي قيمة المستخلص (شامل الضريبة) مطلوب.")
    if min(v["materials"], v["penalties"]) < 0:
        errors.append("الخصومات لا تكون سالبة.")
    if v["materials"] + v["penalties"] > v["value"] + 0.005:
        errors.append("الخصومات (%s) أكبر من قيمة المستخلص (%s)." % ("{:,.2f}".format(v["materials"] + v["penalties"]), "{:,.2f}".format(v["value"])))
    v["date"] = _date("cert_date")
    try:                                                         # what is dated before the opening balance is already inside it
        opening_service.guard_sub_date(sub_id, v["date"], "هذا المستخلص")
    except opening_service.OpeningError as e:
        errors.append(str(e))
    nxt = 1
    v["no"] = int(existing["CertificateNo"]) if existing and existing.get("CertificateNo") else 1
    if v["project"] and v["site"]:
        taken = [g for g in cs.load()[2] if g["SubcontractorID"] == sub_id and norm_key(g.get("Project")) == norm_key(v["project"])
                 and norm_key(g.get("SiteName")) == norm_key(v["site"]) and (not existing or g["CertificateID"] != existing["CertificateID"])]
        nxt = max([g["CertificateNo"] or 0 for g in taken] or [0]) + 1
        try:
            v["no"] = int(f.get("cert_no") or (existing["CertificateNo"] if existing else nxt))
        except ValueError:
            v["no"] = nxt
            errors.append("رقم المستخلص يجب أن يكون رقمًا.")
        if v["no"] <= 0:
            errors.append("رقم المستخلص يجب أن يكون أكبر من صفر.")
        if any(g["CertificateNo"] == v["no"] for g in taken):
            errors.append("رقم المستخلص %s مستخدم في موقع «%s» - الرقم التالي هو %s." % (v["no"], v["site"], nxt))
    files = [x for name in ("final_file", "approved_pdf") for x in request.files.getlist(name) if x and x.filename]
    bad = next((x.filename for x in files if archive._ext(x.filename) not in archive.ALLOWED_EXT), None)
    if bad:
        errors.append("صيغة الملف غير مدعومة: %s (المسموح: PDF، صور، Word، Excel)." % bad)
    return v, errors


def _store_cert_files(cert_id):
    n, user = 0, _user()
    for field, doc_type in (("final_file", "ملف المستخلص من الموقع"), ("approved_pdf", "المستخلص المعتمد")):
        for f in request.files.getlist(field):
            if f and f.filename and save_entity_document("CERT", cert_id, doc_type, f):
                n += 1
    return n


@route("/subcontractors/<int:sub_id>/certificates", methods=["POST"])
@auth.login_required
def certificate_add(sub_id):
    import routes.subcontractors as subs_routes
    if not db.query("SELECT [SubcontractorID] FROM [Subcontractors] WHERE [SubcontractorID] = ?", (sub_id,)):
        flash("لم يتم العثور على المقاول.")
        return redirect(url_for("subcontractors_list"))
    v, errors = _read_cert_form(sub_id)
    if errors:
        for e in errors:
            flash(e)
        return subs_routes.render_sub_page(sub_id, open_dialog="cert-new", form=request.form)
    try:
        with db.transaction() as tx:                            # the project / site records and the certificate - or nothing
            pid, project, project_new = masterdata.resolve_project(v["project"], tx)
            sid, site, site_new = masterdata.resolve_site(pid, v["site"], tx)
            cert_id = tx.execute_insert(
                "INSERT INTO [Certificates] ([SubcontractorID], [ProjectID], [SiteID], [ContractID], [SiteName], [CertificateNo], [CertificateDate], [GrossAmount], [TotalInclVat], "
                "[SuppliedMaterials], [Penalties], [RetentionPercent], [AmountModel], [Status], [CreatedAt]) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, 'v2', ?, ?)",
                (sub_id, pid, sid, v["contract_id"], site, v["no"], v["date"], round(v["value"] / 1.15, 2), v["value"], v["materials"], v["penalties"], ST_ENTERED, datetime.now()))
    except Exception as e:                                 # noqa
        logging.getLogger("po.error").exception("certificate save failed")
        flash("تعذّر حفظ المستخلص ولم يتغير أي شيء. السبب: %s" % " ".join(str(e).split())[:200])
        return subs_routes.render_sub_page(sub_id, open_dialog="cert-new", form=request.form)
    if project_new:
        flash("أُضيف مشروع جديد إلى القائمة: %s" % project)
    if site_new:
        flash("أُضيف موقع جديد للمشروع: %s" % site)
    archive.ensure_cert_folder(cert_id)
    saved = _store_cert_files(cert_id)
    flash("تم حفظ المستخلص رقم %s (%s) وأرشفته%s." % (v["no"], site, (" - رُفع %d ملف" % saved) if saved else " - بلا ملف بعد (يمكن رفعه لاحقًا)"))
    return _back_to_sub(sub_id)


@route("/certificates/<int:cert_id>/edit", methods=["POST"])
@auth.login_required
def certificate_edit(cert_id):
    import routes.subcontractors as subs_routes
    g, contracts, _pays = _cert(cert_id)
    if not g:
        flash("لم يتم العثور على المستخلص.")
        return redirect(url_for("subcontractors_list"))
    sub_id = g["SubcontractorID"]
    if g["Status"] not in (ST_ENTERED, ST_REJECTED) or g.get("StatusUnknown"):
        flash("لا يُعدَّل مستخلص بعد إدراجه في دورة صرف أو صرفه.")
        return _back_to_sub(sub_id)
    try:
        mg.check_can_change(db, cert_id, "التعديل")
    except mg.MoneyError as e:
        flash(str(e))
        return _back_to_sub(sub_id)
    v, errors = _read_cert_form(sub_id, existing=g)
    if errors:
        for e in errors:
            flash(e)
        return subs_routes.render_sub_page(sub_id, open_dialog="cert-%d" % cert_id, form=request.form)
    db.execute("UPDATE [Certificates] SET [CertificateNo] = ?, [CertificateDate] = ?, [GrossAmount] = ?, [TotalInclVat] = ?, [SuppliedMaterials] = ?, "
               "[Penalties] = ?, [RetentionPercent] = 0, [AmountModel] = 'v2', [ContractID] = ? WHERE [CertificateID] = ?",
               (v["no"], v["date"], round(v["value"] / 1.15, 2), v["value"], v["materials"], v["penalties"], v["contract_id"], cert_id))
    saved = _store_cert_files(cert_id)
    flash("تم حفظ تعديلات المستخلص%s." % (" ورُفع %d ملف" % saved if saved else ""))
    return _back_to_sub(sub_id)


@route("/certificates/<int:cert_id>/delete", methods=["POST"])
@auth.login_required
def certificate_delete(cert_id):
    """Only a certificate that has not entered the money flow can be deleted; its folder goes to the trash area."""
    g, _contracts, pays = _cert(cert_id)
    if not g:
        flash("لم يتم العثور على المستخلص.")
        return redirect(url_for("subcontractors_list"))
    sub_id = g["SubcontractorID"]
    if g["Status"] not in (ST_ENTERED, ST_REJECTED) or g.get("CycleID") or any(p["CertificateID"] == cert_id for p in pays):
        flash("لا يُحذف مستخلص أُدرج في دورة صرف أو سُجّل له صرف (تاريخه المالي محفوظ).")
        return _back_to_sub(sub_id)
    try:
        mg.check_can_change(db, cert_id, "الحذف")
    except mg.MoneyError as e:
        flash(str(e))
        return _back_to_sub(sub_id)
    folder = db.query("SELECT [ArchiveFolder] FROM [Certificates] WHERE [CertificateID] = ?", (cert_id,))[0].get("ArchiveFolder")
    with db.transaction() as tx:
        tx.execute("DELETE FROM [CycleDecisions] WHERE [CertificateID] = ?", (cert_id,))
        tx.execute("DELETE FROM [Certificates] WHERE [CertificateID] = ?", (cert_id,))
    trashed = archive.trash_folder(folder) if folder else None
    archive.forget_documents("CERT", cert_id)
    flash("حُذف المستخلص%s." % (" ونُقلت ملفاته إلى سلة المحذوفات في الأرشيف" if trashed else ""))
    return _back_to_sub(sub_id)


@route("/contracts/<int:contract_id>/certificates/new", methods=["GET", "POST"])
@auth.login_required
def certificate_new(contract_id):
    """Old address: certificates are now added from the contractor's page."""
    rows = db.query("SELECT [SubcontractorID] FROM [Contracts] WHERE [ContractID] = ?", (contract_id,))
    return redirect(url_for("subcontractor_detail", sub_id=rows[0]["SubcontractorID"], new_cert=1) if rows else url_for("subcontractors_list"))


@route("/certificates/new", methods=["GET", "POST"])
@auth.login_required
def certificate_new_any():
    """Old address: the certificates tab is follow-up only; adding happens from the contractor's page."""
    cid = request.values.get("contract_id", type=int)
    rows = db.query("SELECT [SubcontractorID] FROM [Contracts] WHERE [ContractID] = ?", (cid,)) if cid else []
    flash("المستخلصات تُضاف من صفحة المقاول.")
    return redirect(url_for("subcontractor_detail", sub_id=rows[0]["SubcontractorID"], new_cert=1) if rows else url_for("subcontractors_list"))


# --------------------------------------------------------------------------- one certificate
def _cert(cert_id):
    subs, contracts, certs, pays = cs.load()
    g = next((x for x in certs if x["CertificateID"] == cert_id), None)
    return g, contracts, pays


@route("/certificates/<int:cert_id>")
@auth.login_required
def certificate_detail(cert_id):
    g, contracts, pays = _cert(cert_id)
    if not g:
        flash("لم يتم العثور على المستخلص.")
        return redirect(url_for("certificates_list"))
    folder = g.get("ArchiveFolder") or archive.cert_folder_rel(g["SubcontractorName"], g["Project"], g.get("SiteName"), g["CertificateNo"])
    return render_template(
        "certificate_detail.html", cert=g, contract=contracts.get(g["ContractID"]), badge=STATUS_BADGE,
        payments=sorted((p for p in pays if p["CertificateID"] == cert_id), key=lambda p: p.get("PaidAt") or datetime.min),
        docs=archive.panel("CERT", cert_id, archive.CERT_DOC_TYPES), folder=folder, folder_exists=bool(g.get("ArchiveFolder")),
        network=archive.network_path(folder), today_iso=datetime.now().strftime("%Y-%m-%d"),
        ST_ENTERED=ST_ENTERED, ST_IN_CYCLE=ST_IN_CYCLE, ST_PAID=ST_PAID, ST_REJECTED=ST_REJECTED)


def _can_change_or_flash(cert_id, what):
    try:
        mg.check_can_change(db, cert_id, what)
        return True
    except mg.MoneyError as e:
        flash(str(e))
        return False


@route("/certificates/<int:cert_id>/reject", methods=["POST"])
@auth.login_required
def certificate_reject(cert_id):
    rows = db.query("SELECT [Status] FROM [Certificates] WHERE [CertificateID] = ?", (cert_id,))
    reason = request.form.get("reason", "").strip()
    if not rows:
        abort(404)
    if norm_status(rows[0]["Status"]) != ST_ENTERED:
        flash("يُرفض المستخلص وهو «مُدخَل» فقط.")
    elif not reason:
        flash("سبب الرفض مطلوب.")
    elif not _can_change_or_flash(cert_id, "الرفض"):
        pass
    else:
        db.execute("UPDATE [Certificates] SET [Status] = ?, [RejectionReason] = ? WHERE [CertificateID] = ?", (ST_REJECTED, reason, cert_id))
        flash("أُعيد المستخلص (مرفوض) مع السبب.")
    return redirect(url_for("certificate_detail", cert_id=cert_id))


@route("/certificates/<int:cert_id>/reopen", methods=["POST"])
@auth.login_required
def certificate_reopen(cert_id):
    rows = db.query("SELECT [Status] FROM [Certificates] WHERE [CertificateID] = ?", (cert_id,))
    if rows and norm_status(rows[0]["Status"]) == ST_REJECTED:
        db.execute("UPDATE [Certificates] SET [Status] = ? WHERE [CertificateID] = ?", (ST_ENTERED, cert_id))
        flash("عاد المستخلص إلى «مُدخَل».")
    return redirect(url_for("certificate_detail", cert_id=cert_id))


@route("/certificates/<int:cert_id>/pay", methods=["POST"])
@auth.login_required
def certificate_pay(cert_id):
    """Records the actual transfer (date, bank reference, voucher). Until the payment-cycle screens exist this is the
    direct way; the payment goes to the payments tab either way."""
    g, contracts, pays = _cert(cert_id)
    if not g:
        abort(404)
    if g["Status"] == ST_IN_CYCLE and g.get("CycleID"):
        flash("هذا المستخلص ضمن دورة صرف: حدّد «تم التحويل / لم يتم التنفيذ» من صفحة الدورة في المرحلة النهائية.")
        return redirect(url_for("cycle_detail", cycle_id=g["CycleID"]))
    if g["Status"] not in (ST_ENTERED, ST_IN_CYCLE) or g.get("StatusUnknown"):
        flash("لا يُسجَّل صرف لمستخلص حالته «%s»." % (g["StatusRaw"] if g.get("StatusUnknown") else g["Status"]))
        return redirect(url_for("certificate_detail", cert_id=cert_id))
    try:
        mg.check_new_payment(db, cert_id, g["Remaining"])
    except mg.MoneyError as e:
        flash(str(e))
        return redirect(url_for("certificate_detail", cert_id=cert_id))
    proof = request.files.get("proof_file")
    if proof and proof.filename and archive._ext(proof.filename) not in archive.ALLOWED_EXT:
        flash("صيغة سند التحويل غير مدعومة - لم يُسجَّل الصرف.")
        return redirect(url_for("certificate_detail", cert_id=cert_id))
    doc = save_entity_document("CERT", cert_id, "سند الصرف", proof, request.form.get("note", "").strip()) if proof and proof.filename else None
    paid_at = _date("pay_date")
    try:
        with db.transaction() as tx:                       # the payment record and the status change together
            mg.check_new_payment(tx, cert_id, g["Remaining"])
            tx.execute("INSERT INTO [Payments] ([CertificateID], [SubcontractorID], [PaidAt], [Amount], [BankRef], [CycleID], [DocID], [Note], "
                       "[CreatedBy], [CreatedAt]) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                       (cert_id, g["SubcontractorID"], paid_at, g["Remaining"], request.form.get("bank_ref", "").strip()[:80], g.get("CycleID"),
                        doc["DocID"] if doc else None, request.form.get("note", "").strip()[:250], _user(), datetime.now()))
            tx.execute("UPDATE [Certificates] SET [Status] = ?, [PaymentDate] = ? WHERE [CertificateID] = ?", (ST_PAID, paid_at, cert_id))
    except mg.MoneyError as e:
        if doc:
            archive.discard_document(doc["DocID"])
        flash(str(e))
        return redirect(url_for("certificate_detail", cert_id=cert_id))
    except Exception as e:                                 # noqa
        if doc:
            archive.discard_document(doc["DocID"])
        logging.getLogger("po.error").exception("payment save failed")
        flash("تعذّر تسجيل الصرف ولم يتغير شيء. السبب: %s" % " ".join(str(e).split())[:200])
        return redirect(url_for("certificate_detail", cert_id=cert_id))
    flash("تم تسجيل صرف %s ريال." % "{:,.2f}".format(g["Remaining"]) + ("" if doc else " (لم يُرفق سند التحويل - يمكن إضافته من المستندات)"))
    return redirect(url_for("certificate_detail", cert_id=cert_id))


@route("/certificates/<int:cert_id>/unpay", methods=["POST"])
@auth.login_required
def certificate_unpay(cert_id):
    """Corrects a payment entered by mistake. The voucher stays in the archive."""
    rows = db.query("SELECT [Status], [CycleID], [ClosedAt] FROM [Certificates] WHERE [CertificateID] = ?", (cert_id,))
    if rows and rows[0].get("ClosedAt"):
        flash("معاملة هذا المستخلص مُقفلة في الأرشيف؛ لا يُلغى صرفها. (إعادة فتحها بصلاحية أعلى أولًا.)")
        return redirect(url_for("certificate_detail", cert_id=cert_id))
    if db.query("SELECT [PaymentID] FROM [Payments] WHERE [CertificateID] = ?", (cert_id,)) and any(p.get("CycleID") for p in db.query("SELECT [CycleID] FROM [Payments] WHERE [CertificateID] = ?", (cert_id,))):
        flash("صُرف هذا المستخلص عبر دورة صرف؛ يُصحَّح من شاشة التحويلات في الدورة.")
        return redirect(url_for("certificate_detail", cert_id=cert_id))
    if rows and classify_status(rows[0]["Status"]) == ST_PAID:
        back_to = ST_IN_CYCLE if rows[0].get("CycleID") else ST_ENTERED
        with db.transaction() as tx:
            tx.execute("DELETE FROM [Payments] WHERE [CertificateID] = ?", (cert_id,))
            tx.execute("UPDATE [Certificates] SET [Status] = ?, [PaymentDate] = NULL WHERE [CertificateID] = ?", (back_to, cert_id))
        flash("أُلغي تسجيل الصرف وعاد المستخلص إلى «%s»." % back_to)
    return redirect(url_for("certificate_detail", cert_id=cert_id))


@route("/certificates/<int:cert_id>/file/<kind>")
@auth.login_required
def certificate_file(cert_id, kind):
    column = {"certificate": "CertificateFile", "payment_proof": "PaymentProofFile"}.get(kind)
    if not column:
        abort(404)
    rows = db.query("SELECT [%s] AS p FROM [Certificates] WHERE [CertificateID] = ?" % column, (cert_id,))
    if not rows or not rows[0]["p"]:
        abort(404)
    return archive.send(rows[0]["p"])
